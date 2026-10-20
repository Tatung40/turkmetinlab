"""Üç (ya da daha çok) insan uzmanlı değerlendirmenin analizi. İki uzmanlı analiz.py'yi değiştirmez; onun dosya
doğrulamasını, erişim ölçütlerini ve bootstrap'ını kullanır.

Önceden belirlenen plan (puanlar görülmeden):

**Uyum**
- Birincil: Krippendorff α (sıralı) ve çok puanlayıcılı, doğrusal ağırlıklı Gwet AC2.
- İkincil:
  - ikili doğrusal ve karesel ağırlıklı Cohen κ;
  - üçlü tam uyum ve ikili ham uyum;
  - karesel ağırlıklı AC2;
  - α (aralık ve ad).
- Güven aralığı: sorgular üzerinden küme bootstrap.

**Uzlaşı**
- Üç dosya da teslim edildikten sonra yapılır.
- Uzlaşı dosyasında, üç puanın aynı olmadığı çiftler yer alır.

**Erişim**
- Birincil etiket: uzlaşı.
- Duyarlılık için: her uzman ayrı ayrı ve ortanca.
- Ölçütler: nDCG@5 ve @10, P@5 ve @10.
- Karşılaştırma: sorgu bazında eşleştirilmiş bootstrap.

**Karar**
- Hibrit − TF-IDF nDCG@5 ≥ +0,030 ve güven aralığının alt sınırı > 0 olmalıdır.
- Pilotta karar verilmez.

    python uzman_dogrulama/coklu_analiz.py --anahtar pilot_insan_anahtar.json \
        --uzman pilot_insan_uzman_1.xlsx pilot_insan_uzman_2.xlsx pilot_insan_uzman_3.xlsx --cikti sonuc
    python uzman_dogrulama/coklu_analiz.py ... --uzlasma sonuc/pilot_insan_uzlasma.xlsx --cikti sonuc2
"""
from __future__ import annotations

import argparse
import itertools
import json
import os
import sys
from collections import Counter

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from uzman_dogrulama import analiz as ia  # noqa: E402
from uzman_dogrulama import ortak as o  # noqa: E402


# ---------------------------------------------------------------------------
# Katsayılar
# ---------------------------------------------------------------------------

def krippendorff_alfa(veri, olcek: str = "sirali") -> float:
    """veri: birim × puanlayıcı (eksik = None/NaN). Krippendorff (2011), "Computing Krippendorff's Alpha-Reliability".

    Çakışma matrisi o_ck; n_c marjinaller; α = 1 − (n − 1) Σ o_ck δ²_ck / Σ n_c n_k δ²_ck.
    δ²: ad 0/1; aralık (c − k)²; sıralı (Σ_{g=c..k} n_g − (n_c + n_k)/2)²; oran ((c − k)/(c + k))².
    """
    birimler = [[float(x) for x in satir if x is not None and not (isinstance(x, float) and np.isnan(x))] for satir in veri]
    birimler = [b for b in birimler if len(b) >= 2]
    degerler = sorted({x for b in birimler for x in b})
    if len(degerler) < 2:
        return float("nan")
    idx = {v: i for i, v in enumerate(degerler)}
    q = len(degerler)
    O = np.zeros((q, q))
    for b in birimler:
        m = len(b)
        for i, j in itertools.permutations(range(m), 2):
            O[idx[b[i]], idx[b[j]]] += 1 / (m - 1)
    nc = O.sum(1)
    n = nc.sum()
    v = np.array(degerler)
    if olcek == "ad":
        D = 1 - np.eye(q)
    elif olcek == "aralik":
        D = (v[:, None] - v[None, :]) ** 2
    elif olcek == "oran":
        D = ((v[:, None] - v[None, :]) / (v[:, None] + v[None, :] + 1e-300)) ** 2
    elif olcek == "sirali":
        D = np.zeros((q, q))
        for c in range(q):
            for k in range(q):
                a, b = min(c, k), max(c, k)
                D[c, k] = (nc[a:b + 1].sum() - (nc[c] + nc[k]) / 2) ** 2
    else:
        raise ValueError(olcek)
    pay = (O * D).sum()
    payda = (np.outer(nc, nc) * D).sum() / (n - 1)
    return float(1 - pay / payda) if payda > 0 else float("nan")


def gwet_ac2_coklu(veri, tur: str = "dogrusal", q: int = 3) -> float:
    """Çok puanlayıcılı Gwet AC1/AC2 (Gwet 2014): r_ik birimde k kategorisini seçen puanlayıcı sayısı,
    r*_ik = Σ_l w_kl r_il; pa = ort_{r_i≥2} Σ_k r_ik (r*_ik − 1) / (r_i (r_i − 1));
    π_k = ort_i r_ik / r_i; pe = T_w / (q(q − 1)) Σ π_k (1 − π_k)."""
    W = o.agirliklar(tur, q)
    R = np.zeros((len(veri), q))
    for i, satir in enumerate(veri):
        for x in satir:
            if x is not None:
                R[i, int(x)] += 1
    ri = R.sum(1)
    cok = ri >= 2
    Ryildiz = R @ W.T
    pa = np.mean((R[cok] * (Ryildiz[cok] - 1)).sum(1) / (ri[cok] * (ri[cok] - 1)))
    pi = (R[ri > 0] / ri[ri > 0, None]).mean(0)
    pe = W.sum() / (q * (q - 1)) * (pi * (1 - pi)).sum()
    return float((pa - pe) / (1 - pe))


# ---------------------------------------------------------------------------
# Raporlar
# ---------------------------------------------------------------------------

def uyum_raporu(anahtar, puanlar: list) -> dict:
    ciftler = [c["cift"] for c in anahtar["ciftler"]]
    sorgu = np.array([c["sorgu_no"] for c in anahtar["ciftler"]])
    X = np.array([[p[c] for p in puanlar] for c in ciftler], float)  # çift × uzman
    k = X.shape[1]
    r = {"cift_sayisi": len(ciftler), "sorgu_sayisi": int(len(set(sorgu))), "uzman_sayisi": k,
         "tam_uyum_yuzde (üçü aynı)": round(100 * float((X.max(1) == X.min(1)).mean()), 2)}
    katsayi = {"alfa_sirali (birincil)": lambda M: krippendorff_alfa(M, "sirali"),
               "ac2_dogrusal (birincil)": lambda M: gwet_ac2_coklu(M, "dogrusal"),
               "ac2_karesel": lambda M: gwet_ac2_coklu(M, "karesel"),
               "alfa_aralik": lambda M: krippendorff_alfa(M, "aralik"), "alfa_ad": lambda M: krippendorff_alfa(M, "ad")}
    for ad, fn in katsayi.items():
        ga = o.kume_bootstrap_uyum(sorgu, X, np.zeros(len(X)), lambda A, _s, fn=fn: fn(A.tolist()), b=1000)
        r[ad] = {"deger": round(fn(X.tolist()), 4), "ga95_sorgu_bootstrap": [round(v, 4) for v in ga]}
    r["ikili"] = {}
    for i, j in itertools.combinations(range(k), 2):
        a, b = X[:, i].astype(int), X[:, j].astype(int)
        r["ikili"][f"uzman_{i + 1}–uzman_{j + 1}"] = {
            "ham_uyum_yuzde": round(100 * float((a == b).mean()), 2),
            "kappa_dogrusal": round(o.cohen_kappa(a, b, "dogrusal"), 4), "kappa_karesel": round(o.cohen_kappa(a, b, "karesel"), 4),
            "capraz_tablo (satır: ilk uzman)": o.capraz_tablo(a, b).astype(int).tolist()}
    r["dagilim"] = {f"uzman_{i + 1}": {o.OLCEK_ADLARI[s]: {"n": int((X[:, i] == s).sum()),
                                                          "yuzde": round(100 * float((X[:, i] == s).mean()), 2)} for s in o.OLCEK}
                    for i in range(k)}
    r["oruntu (sıralı puan üçlüsü)"] = dict(sorted(Counter("-".join(str(int(v)) for v in sorted(x)) for x in X).items()))
    r["aralik (en yüksek − en düşük)"] = {str(d): int(((X.max(1) - X.min(1)) == d).sum()) for d in (0, 1, 2)}
    med = np.median(X, 1)
    r["ortancadan_sapma (uzman başına çift)"] = {f"uzman_{i + 1}": int((X[:, i] != med).sum()) for i in range(k)}
    kaynak = ["+".join(y for y in o.YONTEMLER if y in c["siralar"]) for c in anahtar["ciftler"]]
    uyusmaz = X.max(1) != X.min(1)
    r["uyusmazlik_kaynaga_gore (puanlama sonrası)"] = {
        g: {"cift": int(sum(x == g for x in kaynak)), "uyusmazlik_yuzde": round(100 * float(np.mean(uyusmaz[[x == g for x in kaynak]])), 1)}
        for g in sorted(set(kaynak))}
    return r


UZLASMA_SABIT = ["Çift kimliği", "Sorgu no", "Sorgu başlığı", "Sorgu özeti", "Aday başlığı", "Aday özeti"]


def uzlasma_sutunlari(k: int) -> list:
    return UZLASMA_SABIT + [x for i in range(1, k + 1) for x in (f"Uzman {i} puanı", f"Uzman {i} açıklaması")] + \
        ["Uzlaşı puanı (0/1/2)", "Gerekçe", "Denetim kodu"]


def uzlasma_yaz(yol, anahtar, uzmanlar: list, metinler: dict) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    k = len(uzmanlar)
    sut = uzlasma_sutunlari(k)
    wb = Workbook()
    ws = wb.active
    ws.title = "Uzlaşma"
    ws.append(sut)
    for c in anahtar["ciftler"]:
        p = [u["puan"][c["cift"]] for u in uzmanlar]
        if len({x[0] for x in p}) > 1:
            m = metinler[c["cift"]]
            ws.append([c["cift"], c["sorgu_no"]] + [m[x] for x in UZLASMA_SABIT[2:]] +
                      [v for x in p for v in x] + [None, None, c["denetim"]])
    son = ws.max_row
    uz = sut.index("Uzlaşı puanı (0/1/2)")
    for satir in ws.iter_rows(min_row=2):
        for h in satir:
            h.alignment = Alignment(wrap_text=True, vertical="top")
        for h in satir[uz:uz + 2]:
            h.protection = Protection(locked=False)
    for h in ws[1]:
        h.font = Font(bold=True)
    harf = ws.cell(1, uz + 1).column_letter
    dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True, showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add(f"{harf}2:{harf}{max(son, 2)}")
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(sut)).column_letter}{max(son, 2)}"
    wb.save(yol)
    return son - 1


def uzlasma_oku(yol, anahtar, puanlar: list) -> dict:
    satirlar, _ = ia._sayfa_oku(yol, "Uzlaşma", uzlasma_sutunlari(len(puanlar)))
    beklenen = {c["cift"]: c for c in anahtar["ciftler"] if len({p[c["cift"]] for p in puanlar}) > 1}
    if sorted(r["Çift kimliği"] for r in satirlar) != sorted(beklenen):
        raise o.DogrulamaHatasi("Uzlaşma dosyası uyuşmazlık listesiyle eşleşmiyor.")
    sonuc = {}
    for r in satirlar:
        c = beklenen[r["Çift kimliği"]]
        if r["Denetim kodu"] != c["denetim"] or o.denetim_kodu(c["cift"], str(r["Sorgu özeti"]), str(r["Aday özeti"])) != c["denetim"]:
            raise o.DogrulamaHatasi(f"Uzlaşma: {c['cift']} denetim kodu uyuşmuyor.")
        p = ia._puan(r["Uzlaşı puanı (0/1/2)"])
        if p not in o.OLCEK:
            raise o.DogrulamaHatasi(f"Uzlaşma: {c['cift']} için geçerli uzlaşı puanı yok.")
        sonuc[c["cift"]] = p
    return sonuc


def md(r) -> str:
    u = r["uyum"]
    s = [f"# Üç uzmanlı doğrulama — analiz ({o.ASAMALAR[r['calisma']]})", "", f"Durum: **{r['asama']}**", "",
         "## Uyum", "", f"- Çift: {u['cift_sayisi']} · sorgu: {u['sorgu_sayisi']} · üçlü tam uyum: %{u['tam_uyum_yuzde (üçü aynı)']}"]
    s += [f"- {k}: {v['deger']} {v['ga95_sorgu_bootstrap']}" for k, v in u.items() if isinstance(v, dict) and "deger" in v]
    s += [f"- {k}: ham %{v['ham_uyum_yuzde']}, κ doğrusal {v['kappa_dogrusal']}, κ karesel {v['kappa_karesel']}"
          for k, v in u["ikili"].items()]
    s += [f"- Örüntüler: {u['oruntu (sıralı puan üçlüsü)']}", ""]
    for ad, e in r["erisim"].items():
        s += [f"## Erişim — etiket: {ad}", "", "| Yöntem | nDCG@5 | nDCG@10 | P@5 katı | P@10 katı | P@5 gevşek | P@10 gevşek |",
              "|---|---|---|---|---|---|---|"]
        s += [f"| {y} | " + " | ".join(str(d[m]) for m in ("ndcg@5", "ndcg@10", "p@5_kati", "p@10_kati", "p@5_gevsek", "p@10_gevsek")) + " |"
              for y, d in e["ortalama"].items()]
        s += [f"- {k} nDCG@5: {v['ndcg@5']['fark']} {v['ndcg@5']['ga']}" for k, v in e["farklar (eşleştirilmiş bootstrap, %95 GA)"].items()]
        s.append("")
    s += ["## Karar", "", json.dumps(r["karar"], ensure_ascii=False), "",
          "Sınırlılıklar: küçük örneklem; havuz içi değerlendirme (geri çağırma ölçülemez); 3 düzeyli kısmi ilgililik; "
          "alt alanlar kural tabanlı tahmindir."]
    return "\n".join(s) + "\n"


def ana(argv=None) -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--anahtar", required=True)
    p.add_argument("--uzman", nargs="+", required=True)
    p.add_argument("--uzlasma")
    p.add_argument("--cikti", required=True)
    a = p.parse_args(argv)
    os.makedirs(a.cikti, exist_ok=True)
    anahtar = json.load(open(a.anahtar, encoding="utf-8"))
    calisma = anahtar.get("asama", "asil")
    beklenen_k = int(anahtar.get("olusturma", {}).get("uzman_sayisi", len(a.uzman)))
    on = "pilot_insan_" if calisma.endswith("pilot") else "insan_"

    def hata(e):
        o.json_yaz(os.path.join(a.cikti, f"{on}durum.json"), {"hata": str(e)})
        print("HATA:", e)
        return 2
    try:
        uzmanlar = [ia.uzman_oku(y, anahtar) for y in a.uzman]
        kodlar = [u["kod"] for u in uzmanlar]
        if sorted(kodlar) != list(range(1, beklenen_k + 1)):
            raise o.DogrulamaHatasi(f"Uzman kodları 1..{beklenen_k} olmalı ve her biri bir kez gelmeli; gelen: {kodlar}.")
        uzmanlar.sort(key=lambda u: u["kod"])
        durum = {f"uzman_{u['kod']}": {"eksik": sum(v[0] is None for v in u["puan"].values()),
                                       "gecersiz": sum(v[0] == "geçersiz" for v in u["puan"].values())} for u in uzmanlar}
        if any(v["eksik"] or v["gecersiz"] for v in durum.values()):
            raise o.DogrulamaHatasi(f"Eksik ya da geçersiz puan var: {durum}")
    except o.DogrulamaHatasi as e:
        return hata(e)
    puanlar = [{c: v[0] for c, v in u["puan"].items()} for u in uzmanlar]
    rapor = {"calisma": calisma, "uyum": uyum_raporu(anahtar, puanlar)}
    etiket = {f"uzman_{i + 1}": pu for i, pu in enumerate(puanlar)}
    etiket["ortanca"] = {c: float(np.median([pu[c] for pu in puanlar])) for c in puanlar[0]}
    if a.uzlasma:
        try:
            uz = uzlasma_oku(a.uzlasma, anahtar, puanlar)
        except o.DogrulamaHatasi as e:
            return hata(e)
        etiket = {"uzlasi": {c: uz.get(c, puanlar[0][c]) for c in puanlar[0]}, **etiket}
        rapor["asama"] = "birincil analiz (uzlaşı etiketleri)"
    else:
        metinler, _ = ia._sayfa_oku(a.uzman[0], o.SAYFA, o.SUTUNLAR)
        n = uzlasma_yaz(os.path.join(a.cikti, f"{on}uzlasma.xlsx"), anahtar, uzmanlar,
                        {r["Çift kimliği"]: r for r in metinler})
        rapor["asama"] = f"ön analiz (uzlaşı bekleniyor: {n} çift {on}uzlasma.xlsx dosyasında)"
    rapor["erisim"] = ia.erisim_raporu(anahtar, etiket)
    rapor["karar"] = ({"durum": "PİLOT: karar verilmez; pilot sonuçları asıl analize katılmaz."}
                      if calisma.endswith("pilot") else ia.karar(rapor["erisim"], "asil"))
    o.json_yaz(os.path.join(a.cikti, f"{on}analiz.json"), rapor)
    open(os.path.join(a.cikti, f"{on}analiz.md"), "w", encoding="utf-8").write(md(rapor))
    print(md(rapor))
    return 0


if __name__ == "__main__":
    sys.exit(ana())