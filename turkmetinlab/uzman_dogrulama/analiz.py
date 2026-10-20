"""Uzman puanlarının analizi: dosya doğrulama, puanlayıcılar arası uyum, uzlaşı dosyası ve bilgi erişim ölçütleri.

Yerel çalışır (ağ yok). Önceden belirlenen plan: uzman_dogrulama/ortak.py başındaki açıklama.

    # 1) İki bağımsız dosya geldikten sonra: uyum + uzlaşı dosyası (uzlasma.xlsx) + ön analiz
    python uzman_dogrulama/analiz.py --anahtar anahtar.json --uzman1 uzman_1.xlsx --uzman2 uzman_2.xlsx --cikti sonuc
    # 2) Uzlaşı toplantısından sonra: birincil analiz ve karar
    python uzman_dogrulama/analiz.py ... --uzlasma sonuc/uzlasma.xlsx --cikti sonuc
Çıkış kodu: 0 tamam, 2 eksik/uyumsuz dosya (ayrıntı: durum.json).
"""
from __future__ import annotations

import argparse
import json
import os
import sys
from collections import Counter

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from uzman_dogrulama import ortak as o  # noqa: E402

UZLASMA_SUTUNLARI = ["Çift kimliği", "Sorgu no", "Sorgu başlığı", "Sorgu özeti", "Aday başlığı", "Aday özeti",
                     "Uzman 1 puanı", "Uzman 1 açıklaması", "Uzman 2 puanı", "Uzman 2 açıklaması",
                     "Uzlaşı puanı (0/1/2)", "Gerekçe", "Denetim kodu"]


def _puan(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        f = float(str(v).strip().replace(",", "."))
    except ValueError:
        return "geçersiz"
    return int(f) if f in (0, 1, 2) else "geçersiz"


def _sayfa_oku(yol, sayfa, sutunlar):
    from openpyxl import load_workbook
    wb = load_workbook(yol, read_only=True, data_only=True)
    if sayfa not in wb.sheetnames:
        raise o.DogrulamaHatasi(f"{os.path.basename(yol)}: '{sayfa}' sayfası yok.")
    satirlar = list(wb[sayfa].iter_rows(values_only=True))
    if not satirlar or list(satirlar[0][:len(sutunlar)]) != sutunlar:
        raise o.DogrulamaHatasi(f"{os.path.basename(yol)}: sütun başlıkları değiştirilmiş.")
    ust = [str(r[0]) for r in wb[o.YONERGE_SAYFASI].iter_rows(max_row=8, values_only=True)] \
        if o.YONERGE_SAYFASI in wb.sheetnames else []
    wb.close()
    return [dict(zip(sutunlar, r)) for r in satirlar[1:] if any(x is not None for x in r)], ust


def uzman_oku(yol, anahtar) -> dict:
    """{çift: (puan|None|'geçersiz', açıklama)} — kimlik, havuz ve denetim kodu doğrulanır."""
    satirlar, ust = _sayfa_oku(yol, o.SAYFA, o.SUTUNLAR)
    alan = {x.split(":", 1)[0].strip(): x.split(":", 1)[1].strip() for x in ust if ":" in x}
    try:
        kod, havuz, asama = int(alan["Uzman kodu"]), alan["Havuz kimliği"], alan["Aşama"]
    except (KeyError, ValueError):
        raise o.DogrulamaHatasi(f"{os.path.basename(yol)}: Yönerge sayfasındaki uzman kodu/havuz kimliği/aşama okunamadı.")
    if asama != o.ASAMALAR[anahtar.get("asama", "asil")]:
        raise o.DogrulamaHatasi(f"{os.path.basename(yol)}: aşama uyuşmuyor (pilot ve asıl çalışma karıştırılamaz).")
    if havuz != anahtar["havuz_kimligi"]:
        raise o.DogrulamaHatasi(f"{os.path.basename(yol)}: başka bir havuza ait (havuz kimliği uyuşmuyor).")
    beklenen = {c["cift"]: c for c in anahtar["ciftler"]}
    sayac = Counter(r["Çift kimliği"] for r in satirlar)
    tekrar = [k for k, n in sayac.items() if n > 1]
    if tekrar:
        raise o.DogrulamaHatasi(f"{os.path.basename(yol)}: yinelenen çift kimliği: {tekrar[:5]}")
    fazla, eksik = set(sayac) - set(beklenen), set(beklenen) - set(sayac)
    if fazla or eksik:
        raise o.DogrulamaHatasi(f"{os.path.basename(yol)}: anahtarla eşleşmeyen satırlar (fazla {len(fazla)}, "
                                f"eksik {len(eksik)}).")
    sonuc = {}
    for r in satirlar:
        c = beklenen[r["Çift kimliği"]]
        if o.denetim_kodu(c["cift"], str(r["Sorgu özeti"]), str(r["Aday özeti"])) != c["denetim"] or \
                r["Denetim kodu"] != c["denetim"] or r["Sorgu no"] != c["sorgu_no"]:
            raise o.DogrulamaHatasi(f"{os.path.basename(yol)}: {c['cift']} satırının metni/denetim kodu anahtarla "
                                    "uyuşmuyor (satır kaydırılmış ya da düzenlenmiş olabilir).")
        sonuc[c["cift"]] = (_puan(r["Puan (0/1/2)"]), r["Açıklama (isteğe bağlı)"])
    return {"kod": kod, "puan": sonuc}


def uyum_raporu(anahtar, p1, p2):
    ciftler = [c["cift"] for c in anahtar["ciftler"]]
    sorgu = np.array([c["sorgu_no"] for c in anahtar["ciftler"]])
    a, b = np.array([p1[c] for c in ciftler]), np.array([p2[c] for c in ciftler])
    n = len(a)
    r = {"cift_sayisi": n, "sorgu_sayisi": int(len(set(sorgu))), "ham_uyum_yuzde": round(100 * float((a == b).mean()), 2)}
    for ad, fn, tur in (("kappa_dogrusal (birincil)", o.cohen_kappa, "dogrusal"), ("ac2_dogrusal (birincil)", o.gwet_ac, "dogrusal"),
                        ("kappa_karesel", o.cohen_kappa, "karesel"), ("ac2_karesel", o.gwet_ac, "karesel"),
                        ("kappa_agirliksiz", o.cohen_kappa, "yok"), ("ac1", o.gwet_ac, "yok")):
        f = lambda x, y, fn=fn, tur=tur: fn(x, y, tur)  # noqa: E731
        r[ad] = {"deger": round(f(a, b), 4), "ga95_sorgu_bootstrap": [round(v, 4) for v in o.kume_bootstrap_uyum(sorgu, a, b, f)]}
    T = o.capraz_tablo(a, b)
    r["capraz_tablo (satır: uzman 1, sütun: uzman 2)"] = T.astype(int).tolist()
    r["dagilim"] = {f"uzman_{i}": {o.OLCEK_ADLARI[k]: {"n": int((x == k).sum()), "yuzde": round(100 * float((x == k).mean()), 2)}
                                   for k in o.OLCEK} for i, x in ((1, a), (2, b))}
    fark = a - b
    r["uyusmazlik"] = {"toplam": int((fark != 0).sum()), "bir_basamak": int((np.abs(fark) == 1).sum()),
                       "iki_basamak (0↔2)": int((np.abs(fark) == 2).sum()), "uzman_1_yuksek": int((fark > 0).sum()),
                       "uzman_2_yuksek": int((fark < 0).sum())}
    kaynak = ["+".join(y for y in o.YONTEMLER if y in c["siralar"]) for c in anahtar["ciftler"]]
    r["uyusmazlik_kaynaga_gore (puanlama sonrası, anahtardan)"] = {
        k: {"cift": int(sum(x == k for x in kaynak)), "uyusmazlik_yuzde": round(100 * float(np.mean([f != 0 for f, x in zip(fark, kaynak) if x == k])), 1)}
        for k in sorted(set(kaynak))}
    sorgu_oran = sorted(((round(float((fark[sorgu == s] != 0).mean()), 3), int(s)) for s in set(sorgu)), reverse=True)
    r["en_cok_uyusmazlik_olan_sorgular"] = [{"sorgu_no": s, "oran": v} for v, s in sorgu_oran[:5]]
    return r


def uzlasma_yaz(yol, anahtar, p1, p2, metinler):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Uzlaşma"
    ws.append(UZLASMA_SUTUNLARI)
    for c in anahtar["ciftler"]:
        k = c["cift"]
        if p1[k][0] != p2[k][0]:
            m = metinler[k]
            ws.append([k, c["sorgu_no"], m["Sorgu başlığı"], m["Sorgu özeti"], m["Aday başlığı"], m["Aday özeti"],
                       p1[k][0], p1[k][1], p2[k][0], p2[k][1], None, None, c["denetim"]])
    son = ws.max_row
    for satir in ws.iter_rows(min_row=2):
        for h in satir:
            h.alignment = Alignment(wrap_text=True, vertical="top")
        for h in satir[10:12]:
            h.protection = Protection(locked=False)
    for h in ws[1]:
        h.font = Font(bold=True)
    for i, g in zip("ABCDEFGHIJKLM", (11, 8, 28, 50, 28, 50, 9, 22, 9, 22, 12, 30, 12)):
        ws.column_dimensions[i].width = g
    dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True, showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add(f"K2:K{max(son, 2)}")
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.auto_filter.ref = f"A1:M{max(son, 2)}"
    wb.save(yol)
    return son - 1


def uzlasma_oku(yol, anahtar, p1, p2) -> dict:
    satirlar, _ = _sayfa_oku(yol, "Uzlaşma", UZLASMA_SUTUNLARI)
    beklenen = {c["cift"]: c for c in anahtar["ciftler"] if p1[c["cift"]][0] != p2[c["cift"]][0]}
    gelen = [r["Çift kimliği"] for r in satirlar]
    if sorted(gelen) != sorted(beklenen):
        raise o.DogrulamaHatasi("Uzlaşma dosyası uyuşmazlık listesiyle eşleşmiyor.")
    sonuc = {}
    for r in satirlar:
        c = beklenen[r["Çift kimliği"]]
        if r["Denetim kodu"] != c["denetim"] or o.denetim_kodu(c["cift"], str(r["Sorgu özeti"]), str(r["Aday özeti"])) != c["denetim"]:
            raise o.DogrulamaHatasi(f"Uzlaşma: {c['cift']} denetim kodu uyuşmuyor.")
        p = _puan(r["Uzlaşı puanı (0/1/2)"])
        if p not in o.OLCEK:
            raise o.DogrulamaHatasi(f"Uzlaşma: {c['cift']} için geçerli uzlaşı puanı yok.")
        sonuc[c["cift"]] = p
    return sonuc


def erisim_raporu(anahtar, etiketler: dict) -> dict:
    """etiketler: {etiket_kumesi: {çift: ilgililik}}. Her yöntem ve sorgu için ölçütler + eşleştirilmiş farklar."""
    cift_bilgi = {c["cift"]: c for c in anahtar["ciftler"]}
    sorgular = sorted(anahtar["siralamalar"], key=int)
    rapor = {}
    for ad, et in etiketler.items():
        olcut = {y: {m: [] for m in ("ndcg@5", "ndcg@10", "p@5_kati", "p@10_kati", "p@5_gevsek", "p@10_gevsek",
                                     "havuz_ici_geri_cagirma@10_gevsek")} for y in o.YONTEMLER}
        ilgisiz_sorgu = 0
        for s in sorgular:
            havuz = {c: et[c] for c, b in cift_bilgi.items() if str(b["sorgu_no"]) == s}
            ilgisiz_sorgu += max(havuz.values()) == 0
            for y in o.YONTEMLER:
                sira = anahtar["siralamalar"][s][y]
                for k in (5, 10):
                    olcut[y][f"ndcg@{k}"].append(o.ndcg(sira, havuz, k))
                    olcut[y][f"p@{k}_kati"].append(o.precision(sira, havuz, k, 2))
                    olcut[y][f"p@{k}_gevsek"].append(o.precision(sira, havuz, k, 1))
                ilgili = {c for c, v in havuz.items() if v >= 1}  # havuz yanlılığı göstergesi (gerçek geri çağırma değil)
                olcut[y]["havuz_ici_geri_cagirma@10_gevsek"].append(len(ilgili & set(sira[:10])) / len(ilgili)
                                                                     if ilgili else np.nan)
        olcut = {y: {m: np.array(v, float) for m, v in d.items()} for y, d in olcut.items()}
        r = {"ilgili_belgesi_olmayan_sorgu (nDCG tanımsız, dışlandı)": int(ilgisiz_sorgu),
             "yargilanmis_kapsam_ilk10": 1.0,
             "ortalama": {y: {m: round(float(np.nanmean(v)), 4) for m, v in d.items()} for y, d in olcut.items()},
             "farklar (eşleştirilmiş bootstrap, %95 GA)": {}}
        for x, y in (("E5", "TF-IDF"), ("Hibrit", "TF-IDF"), ("Hibrit", "E5")):
            r["farklar (eşleştirilmiş bootstrap, %95 GA)"][f"{x} − {y}"] = {
                m: {k: (round(v, 4) if isinstance(v, float) else [round(z, 4) for z in v] if isinstance(v, list) else v)
                    for k, v in o.eslestirilmis_bootstrap(olcut[x][m], olcut[y][m]).items()} for m in olcut[x]}
        rapor[ad] = r
    return rapor


def karar(erisim: dict, asama: str = "asil") -> dict:
    if asama == "pilot":
        return {"durum": "PİLOT: karar verilmez. Pilot yalnızca yönerge ve süreci denemek içindir; asıl analize katılmaz."}
    if "uzlasi" not in erisim:
        return {"durum": "Karar verilmez: uzlaşı etiketleri henüz yok (yalnızca ön analiz)."}
    f = erisim["uzlasi"]["farklar (eşleştirilmiş bootstrap, %95 GA)"]["Hibrit − TF-IDF"]["ndcg@5"]
    var = f["fark"] >= o.PRATIK_ESIK and f["ga"][0] > 0
    return {"olcut": "Hibrit − TF-IDF nDCG@5 ≥ +0,030 ve %95 GA alt sınırı > 0 (önceden belirlendi)", "fark": f["fark"],
            "ga": f["ga"], "pratik_onem": bool(var)}


def md(r) -> str:
    s = [f"# Uzman doğrulaması — analiz ({o.ASAMALAR[r['calisma']]})", "", f"Durum: **{r['asama']}**", "",
         "## Puanlayıcılar arası uyum", ""]
    u = r["uyum"]
    s.append(f"- Çift: {u['cift_sayisi']} · sorgu: {u['sorgu_sayisi']} · ham uyum: %{u['ham_uyum_yuzde']}")
    for k, v in u.items():
        if isinstance(v, dict) and "deger" in v:
            s.append(f"- {k}: {v['deger']} [{v['ga95_sorgu_bootstrap'][0]}, {v['ga95_sorgu_bootstrap'][1]}]")
    s += [f"- Uyuşmazlık: {u['uyusmazlik']}", f"- Dağılım: {json.dumps(u['dagilim'], ensure_ascii=False)}", ""]
    for ad, e in r["erisim"].items():
        s += [f"## Bilgi erişim ölçütleri — etiket: {ad}", "", "| Yöntem | nDCG@5 | nDCG@10 | P@5 katı | P@10 katı | P@5 gevşek | P@10 gevşek |",
              "|---|---|---|---|---|---|---|"]
        for y, d in e["ortalama"].items():
            s.append(f"| {y} | " + " | ".join(str(d[m]) for m in ("ndcg@5", "ndcg@10", "p@5_kati", "p@10_kati", "p@5_gevsek", "p@10_gevsek")) + " |")
        for k, v in e["farklar (eşleştirilmiş bootstrap, %95 GA)"].items():
            s.append(f"- {k} nDCG@5: {v['ndcg@5']['fark']} {v['ndcg@5']['ga']} · nDCG@10: {v['ndcg@10']['fark']} {v['ndcg@10']['ga']}")
        s.append("")
    s += ["## Önceden belirlenen karar", "", json.dumps(r["karar"], ensure_ascii=False), "",
          "Sınırlılıklar: 50 sorgu; değerlendirme yalnızca üç yöntemin ilk 10 sonucundan oluşan havuz içindedir (geri "
          "çağırma ölçülemez); 3 düzeyli kısmi ilgililik; iki uzman; çoğunluğu mühendislik özetleri."]
    return "\n".join(s) + "\n"


def ana(argv=None) -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--anahtar", required=True)
    p.add_argument("--uzman1", required=True)
    p.add_argument("--uzman2", required=True)
    p.add_argument("--uzlasma")
    p.add_argument("--cikti", required=True)
    a = p.parse_args(argv)
    os.makedirs(a.cikti, exist_ok=True)
    anahtar = json.load(open(a.anahtar, encoding="utf-8"))
    calisma = anahtar.get("asama", "asil")
    on = "pilot_" if calisma == "pilot" else ""  # pilot çıktıları asıl çıktılarla aynı adı taşımaz
    try:
        u1, u2 = uzman_oku(a.uzman1, anahtar), uzman_oku(a.uzman2, anahtar)
        if {u1["kod"], u2["kod"]} != {1, 2}:
            raise o.DogrulamaHatasi(f"Uzman kodları 1 ve 2 olmalı; gelen: {u1['kod']}, {u2['kod']} (aynı dosya iki kez?).")
        if u1["kod"] == 2:
            u1, u2 = u2, u1
        durum = {f"uzman_{i}": {"eksik": sum(v[0] is None for v in u["puan"].values()),
                                "gecersiz": sum(v[0] == "geçersiz" for v in u["puan"].values())} for i, u in ((1, u1), (2, u2))}
        if any(v["eksik"] or v["gecersiz"] for v in durum.values()):
            raise o.DogrulamaHatasi(f"Eksik ya da geçersiz puan var: {durum}")
    except o.DogrulamaHatasi as e:
        o.json_yaz(os.path.join(a.cikti, f"{on}durum.json"), {"hata": str(e)})
        print("HATA:", e)
        return 2
    p1 = {k: v[0] for k, v in u1["puan"].items()}
    p2 = {k: v[0] for k, v in u2["puan"].items()}
    rapor = {"calisma": calisma, "uyum": uyum_raporu(anahtar, p1, p2)}
    etiket = {"uzman_1": p1, "uzman_2": p2, "ortalama": {k: (p1[k] + p2[k]) / 2 for k in p1}}
    if a.uzlasma:
        try:
            uz = uzlasma_oku(a.uzlasma, anahtar, u1["puan"], u2["puan"])
        except o.DogrulamaHatasi as e:
            o.json_yaz(os.path.join(a.cikti, f"{on}durum.json"), {"hata": str(e)})
            print("HATA:", e)
            return 2
        etiket = {"uzlasi": {k: (uz[k] if k in uz else p1[k]) for k in p1}, **etiket}
        rapor["asama"] = "birincil analiz (uzlaşı etiketleri)"
    else:
        metinler, _ = _sayfa_oku(a.uzman1, o.SAYFA, o.SUTUNLAR)
        n = uzlasma_yaz(os.path.join(a.cikti, f"{on}uzlasma.xlsx"), anahtar, u1["puan"], u2["puan"],
                        {r["Çift kimliği"]: r for r in metinler})
        rapor["asama"] = f"ön analiz (uzlaşı bekleniyor: {n} uyuşmazlık {on}uzlasma.xlsx dosyasında)"
    rapor["erisim"] = erisim_raporu(anahtar, etiket)
    rapor["karar"] = karar(rapor["erisim"], calisma)
    o.json_yaz(os.path.join(a.cikti, f"{on}analiz.json"), rapor)
    open(os.path.join(a.cikti, f"{on}analiz.md"), "w", encoding="utf-8").write(md(rapor))
    print(md(rapor))
    return 0


if __name__ == "__main__":
    sys.exit(ana())