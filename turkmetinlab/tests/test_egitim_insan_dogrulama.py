"""Eğitim Programları ve Öğretim — üç insan uzmanlı doğrulama: SENTETİK testler (gerçek makale ya da puan yok).

Kapsam: katsayıların yayımlanmış/bağımsız değerlerle denetimi, derlem dahil etme kuralları, sahte OpenAlex sunucusuyla
toplama, yeterlilik kapısı, pilot/asıl ayrımı, üç kör dosya ve üç puanlayıcılı analiz akışı.
"""
import functools
import http.server
import json
import os
import threading
from urllib.parse import parse_qs, urlparse

import numpy as np
import pytest
from openpyxl import load_workbook

from anlamsal_model import kucuk_model_olustur
from uzman_dogrulama import coklu_analiz as ca
from uzman_dogrulama import egitim_derlem as ed
from uzman_dogrulama import egitim_havuz as eh
from uzman_dogrulama import ortak as o

N = None


# --- Katsayılar -----------------------------------------------------------------

def test_krippendorff_yayimlanmis_ornek():
    """Krippendorff (2011), 4 gözlemci × 12 birim, eksik verili örnek: α ad .743, sıralı .815, aralık .849, oran .797."""
    A = [1, 2, 3, 3, 2, 1, 4, 1, 2, N, N, N]
    B = [1, 2, 3, 3, 2, 2, 4, 1, 2, 5, N, 3]
    C = [N, 3, 3, 3, 2, 3, 4, 2, 2, 5, 1, N]
    D = [1, 2, 3, 3, 2, 4, 4, 1, 2, 5, 1, N]
    veri = [list(t) for t in zip(A, B, C, D)]
    beklenen = {"ad": 0.743, "sirali": 0.815, "aralik": 0.849, "oran": 0.797}
    for olcek, deger in beklenen.items():
        assert ca.krippendorff_alfa(veri, olcek) == pytest.approx(deger, abs=5e-4)
    assert ca.krippendorff_alfa([[1, 1, 1], [2, 2, 2], [0, 0, 0]]) == pytest.approx(1)


def test_coklu_ac2_iki_puanlayicida_ikili_formulle_ayni():
    rng = np.random.default_rng(7)
    a = rng.integers(0, 3, 150)
    b = np.where(rng.random(150) < 0.6, a, rng.integers(0, 3, 150))
    for tur in ("yok", "dogrusal", "karesel"):
        assert ca.gwet_ac2_coklu(np.c_[a, b].tolist(), tur) == pytest.approx(o.gwet_ac(a, b, tur), abs=1e-12)
    X = np.c_[a, a, a]
    assert ca.gwet_ac2_coklu(X.tolist()) == pytest.approx(1)
    # Simetri: puanlayıcıların sırası sonucu değiştirmez
    c = rng.integers(0, 3, 150)
    assert ca.gwet_ac2_coklu(np.c_[a, b, c].tolist()) == pytest.approx(ca.gwet_ac2_coklu(np.c_[c, a, b].tolist()))


# --- Derlem kuralları ----------------------------------------------------------------

UZUN = " ".join(["öğrencilerin", "sınıf", "ortamında", "çalışma", "verileri", "analiz", "edilmiştir"] * 9)


def _kayit(baslik, ozet, **k):
    return {"kaynak_kimligi": k.get("kimlik", "W1"), "doi": k.get("doi"), "baslik": baslik, "ozet": ozet,
            "erisim_adresi": k.get("adres", "https://dergipark.org.tr/tr/pub/x/1"), "yil": 2020, "dergi": "D",
            "platform": "DergiPark"}


def test_dahil_etme_kurallari():
    assert ed.ozet_kur({"bir": [0, 3], "üç": [2], "iki": [1]}) == "bir iki üç bir"
    iyi = _kayit("Fen bilimleri öğretim programının değerlendirilmesi", "Bu çalışmada öğretim programı ve kazanım " + UZUN)
    assert ed.dahil_mi(iyi) == (True, "")
    assert not ed.dahil_mi(_kayit("Eğitimde liderlik", "Eğitim yöneticileri ve eğitim " + UZUN))[0]  # yalnızca "eğitim"
    assert "çekirdek" in ed.dahil_mi(_kayit("Kazanım ve hedef", "kazanım hedef içerik " + UZUN))[1]  # destek ifadeler yetmez
    assert "kısa" in ed.dahil_mi(_kayit("Öğretim programı", "öğretim programı kazanım kısa"))[1]
    assert "doğrulanamayan" in ed.dahil_mi({**iyi, "erisim_adresi": ""})[1]
    ing = _kayit("Öğretim programı", "öğretim programı kazanım " + " ".join(["the", "of", "and", "in", "to", "is"] * 15) + UZUN)
    assert "İngilizce" in ed.dahil_mi(ing)[1]
    kopya = [iyi, {**iyi, "kaynak_kimligi": "W2"}, {**iyi, "kaynak_kimligi": "W3", "baslik": "  " + iyi["baslik"] + "."}]
    secilen, neden = ed.suz(kopya)
    assert len(secilen) == 1 and neden["yinelenen kayıt"] == 2
    assert secilen[0]["birincil_alt_alan"] == "Program geliştirme ve değerlendirme"


# --- Sahte OpenAlex sunucusu -------------------------------------------------------------

def _ters_dizin(metin):  # OpenAlex biçimi: sözcük → tüm konumları
    d = {}
    for j, w in enumerate(metin.split()):
        d.setdefault(w, []).append(j)
    return d


class _OpenAlex(http.server.BaseHTTPRequestHandler):
    sayfa_sayisi = 2
    hata_kalan = {"n": 0}

    def log_message(self, *a):
        pass

    def do_GET(self):
        if self.hata_kalan["n"] > 0:
            self.hata_kalan["n"] -= 1
            self.send_response(429)
            self.end_headers()
            return
        q = parse_qs(urlparse(self.path).query)
        imlec = q["cursor"][0]
        ifade = q["search"][0]
        n = 0 if imlec == "*" else int(imlec)
        destek = next(d[0] for c, d in ed.ALT_ALANLAR.values() if ifade.strip('"') in c)
        sonuc = [] if n >= self.sayfa_sayisi else [{
            "id": f"https://openalex.org/W{abs(hash((ifade, n, i))) % 10**8}", "doi": None,
            "display_name": f"Öğretim programı çalışması {ifade} {n} {i}",
            "abstract_inverted_index": _ters_dizin(f"{ifade.strip(chr(34))} {destek} {UZUN} sayfa{n} sira{i}"),
            "publication_year": 2021, "primary_location": {"landing_page_url": f"https://dergipark.org.tr/x/{n}{i}",
                                                           "source": {"display_name": "Sahte Dergi"}}, "topics": []}
            for i in range(3)]
        govde = json.dumps({"results": sonuc, "meta": {"next_cursor": str(n + 1) if sonuc else None}}).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(govde)


@pytest.fixture
def openalex(monkeypatch):
    s = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _OpenAlex)
    threading.Thread(target=s.serve_forever, daemon=True).start()
    monkeypatch.setattr(ed, "ADRES", f"http://127.0.0.1:{s.server_address[1]}")
    monkeypatch.setattr(ed.time, "sleep", lambda *_: None)
    yield
    s.shutdown()


def test_toplama_sayfalama_ve_tek_yeniden_deneme(openalex, tmp_path):
    ham, sayac = ed.topla(["öğretim programı", "öğretim tasarımı"], bekle=0)
    assert len(ham) == 12 and sayac == {"öğretim programı": 6, "öğretim tasarımı": 6}  # 2 sayfa × 3 × 2 ifade
    assert all(k["platform"] == "DergiPark" and k["ozet"].startswith(("öğretim programı", "öğretim tasarımı")) for k in ham)
    assert ed.suz(ham)[0] and all(ed.dahil_mi(k)[0] for k in ham)  # özet ters dizinden doğru kuruldu
    _OpenAlex.hata_kalan["n"] = 1  # bir 429 sonrası devam eder
    assert len(ed.topla(["öğretim programı"], bekle=0)[0]) == 6
    _OpenAlex.hata_kalan["n"] = 2  # iki ardışık 429: durur (sonsuz deneme yok)
    with pytest.raises(Exception):
        ed.topla(["öğretim programı"], bekle=0)
    _OpenAlex.hata_kalan["n"] = 0


# --- Havuz: yeterlilik, pilot/asıl ayrımı, üç kör dosya -------------------------------------

ALAN_METNI = {
    "Program geliştirme ve değerlendirme": "öğretim programı kazanım program değerlendirme",
    "Öğretim tasarımı": "öğretim tasarımı addie öğretim materyali",
    "Öğrenme-öğretme süreçleri": "öğrenme-öğretme süreci öğrenme ortamı sınıf içi",
    "Öğretim stratejileri ve yöntemleri": "öğretim stratejisi deney grubu akademik başarı",
    "Program uygulamaları": "programın uygulanması öğretmen görüşleri karşılaşılan sorunlar"}


def _derlem(yol, n):
    rng = np.random.default_rng(0)
    kayitlar = []
    for i in range(n):
        alan = list(ALAN_METNI)[i % 5]
        ek = " ".join(rng.choice(["öğrenci", "ders", "veri", "nitel", "nicel", "görüş", "sınıf", "okul"], 50))
        kayitlar.append(_kayit(f"Çalışma {i}: {ALAN_METNI[alan].split()[0]} {ALAN_METNI[alan].split()[1]}",
                               f"{ALAN_METNI[alan]} {ek} kayıt{i}", kimlik=f"W{i:05d}", adres=f"https://dergipark.org.tr/a/{i}"))
    secilen, _ = ed.suz(kayitlar)
    with open(yol, "w", encoding="utf-8") as f:
        for k in secilen:
            f.write(json.dumps(k, ensure_ascii=False) + "\n")
    return secilen


@pytest.fixture(scope="module")
def hazir(tmp_path_factory):
    d = tmp_path_factory.mktemp("egitim")
    kucuk_model_olustur(str(d / "model"), azami_token=4000)
    _derlem(d / "derlem.jsonl", 400)
    for asama, kl in (("egitim_pilot", "gp"), ("egitim_asil", "ga")):
        assert eh.ana(["--derlem", str(d / "derlem.jsonl"), "--onnx-klasor", str(d / "model"), "--asama", asama,
                       "--cikti", str(d / kl), "--ozet", str(d / "ozet"), "--n", "50"]) == 0
    return d


def test_yetersiz_derlemde_pilot_uretilmez(tmp_path):
    kucuk_model_olustur(str(tmp_path / "model"), azami_token=4000)
    _derlem(tmp_path / "az.jsonl", 120)
    assert eh.ana(["--derlem", str(tmp_path / "az.jsonl"), "--onnx-klasor", str(tmp_path / "model"), "--asama",
                   "egitim_pilot", "--cikti", str(tmp_path / "g"), "--ozet", str(tmp_path / "oz")]) == 3
    ozet = json.load(open(tmp_path / "oz" / "egitim_havuz_ozeti_egitim_pilot.json", encoding="utf-8"))
    assert "YETERSİZ" in ozet["durum"] and not ozet["yeterlilik"]["yeterli"] and not (tmp_path / "g").exists()


def test_pilot_uc_kor_dosya_ve_asil_ile_ayrim(hazir):
    p = json.load(open(hazir / "gp" / "yalnizca_arastirmaci" / "pilot_insan_anahtar.json", encoding="utf-8"))
    a = json.load(open(hazir / "ga" / "yalnizca_arastirmaci" / "insan_anahtar.json", encoding="utf-8"))
    assert p["olusturma"]["uzman_sayisi"] == 3 and len(p["siralamalar"]) == 5 and len(a["siralamalar"]) == 50
    assert not {c["sorgu_belge"] for c in p["ciftler"]} & {c["sorgu_belge"] for c in a["ciftler"]}
    assert all(c["cift"].startswith("IP-") for c in p["ciftler"]) and p["havuz_kimligi"] != a["havuz_kimligi"]
    assert len({(c["sorgu_belge"], c["aday_belge"]) for c in p["ciftler"]}) == len(p["ciftler"])
    assert all(c["sorgu_belge"] != c["aday_belge"] for c in p["ciftler"])
    assert len(set(p["olusturma"]["sorgular"][s]["alt_alan"] for s in p["olusturma"]["sorgular"])) == 5  # 5 farklı alt alan
    dosyalar = sorted(os.listdir(hazir / "gp" / "uzmanlara_gonderin"))
    assert dosyalar == ["pilot_insan_uzman_1.xlsx", "pilot_insan_uzman_2.xlsx", "pilot_insan_uzman_3.xlsx"]
    siralar = []
    for f in dosyalar:
        wb = load_workbook(hazir / "gp" / "uzmanlara_gonderin" / f)
        ws = wb[o.SAYFA]
        assert wb.sheetnames == [o.YONERGE_SAYFASI, o.SAYFA] and ws.protection.sheet
        assert "EĞİTİM İNSAN PİLOTU" in wb[o.YONERGE_SAYFASI]["A3"].value
        assert ws.cell(2, 1).protection.locked and not ws.cell(2, 7).protection.locked
        for satir in wb[o.SAYFA].iter_rows(values_only=True):
            assert not any(y in str(h) for h in satir if h is not None for y in o.YASAK_IFADELER)
        siralar.append([r[0] for r in list(ws.iter_rows(values_only=True))[1:]])
    assert set(siralar[0]) == set(siralar[1]) == set(siralar[2]) == {c["cift"] for c in p["ciftler"]}
    assert len({tuple(s) for s in siralar}) == 3  # üç farklı karıştırma
    ozet = json.load(open(hazir / "ozet" / "egitim_havuz_ozeti_egitim_pilot.json", encoding="utf-8"))
    assert ozet["durum"] == "üretildi" and "BİLİNMİYOR" in ozet["kaynak_lisans_durumu"]


def _doldur(hazir, u, fn, hedef):
    wb = load_workbook(hazir / "gp" / "uzmanlara_gonderin" / f"pilot_insan_uzman_{u}.xlsx")
    ws = wb[o.SAYFA]
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 7).value = fn(ws.cell(r, 1).value)
    wb.save(hedef)
    return str(hedef)


def test_uc_puanlayici_analizi_uzlasi_ve_hatalar(hazir, tmp_path):
    anahtar = str(hazir / "gp" / "yalnizca_arastirmaci" / "pilot_insan_anahtar.json")
    a = json.load(open(anahtar, encoding="utf-8"))
    hib = {c for s in a["siralamalar"].values() for c in s["Hibrit"][:3]}
    gercek = {c["cift"]: (2 if c["cift"] in hib else 0) for c in a["ciftler"]}
    farkli = sorted(gercek)[:4]
    u = [_doldur(hazir, 1, gercek.get, tmp_path / "1.xlsx"),
         _doldur(hazir, 2, lambda c: 1 if c in farkli else gercek[c], tmp_path / "2.xlsx"),
         _doldur(hazir, 3, gercek.get, tmp_path / "3.xlsx")]
    assert ca.ana(["--anahtar", anahtar, "--uzman", *u, "--cikti", str(tmp_path / "s")]) == 0
    r = json.load(open(tmp_path / "s" / "pilot_insan_analiz.json", encoding="utf-8"))
    assert r["uyum"]["uzman_sayisi"] == 3 and r["uyum"]["aralik (en yüksek − en düşük)"]["0"] == len(gercek) - 4
    assert r["uyum"]["ikili"]["uzman_1–uzman_3"]["kappa_dogrusal"] == pytest.approx(1)
    assert "PİLOT" in r["karar"]["durum"] and r["erisim"]["ortanca"]["ortalama"] == r["erisim"]["uzman_1"]["ortalama"]
    wb = load_workbook(tmp_path / "s" / "pilot_insan_uzlasma.xlsx")
    ws = wb["Uzlaşma"]
    assert ws.max_row - 1 == 4
    sut = ca.uzlasma_sutunlari(3).index("Uzlaşı puanı (0/1/2)") + 1
    for i in range(2, ws.max_row + 1):
        ws.cell(i, sut).value = gercek[ws.cell(i, 1).value]
    wb.save(tmp_path / "uz.xlsx")
    assert ca.ana(["--anahtar", anahtar, "--uzman", *u, "--uzlasma", str(tmp_path / "uz.xlsx"), "--cikti", str(tmp_path / "s2")]) == 0
    r2 = json.load(open(tmp_path / "s2" / "pilot_insan_analiz.json", encoding="utf-8"))
    assert r2["erisim"]["uzlasi"]["ortalama"] == r["erisim"]["uzman_1"]["ortalama"] and "PİLOT" in r2["karar"]["durum"]
    # Hatalar: eksik uzman, aynı dosya iki kez, aşaması farklı anahtar
    assert ca.ana(["--anahtar", anahtar, "--uzman", u[0], u[1], "--cikti", str(tmp_path / "h1")]) == 2
    assert ca.ana(["--anahtar", anahtar, "--uzman", u[0], u[0], u[2], "--cikti", str(tmp_path / "h2")]) == 2
    asil = str(hazir / "ga" / "yalnizca_arastirmaci" / "insan_anahtar.json")
    assert ca.ana(["--anahtar", asil, "--uzman", *u, "--cikti", str(tmp_path / "h3")]) == 2
    assert "aşama" in json.load(open(tmp_path / "h3" / "insan_durum.json", encoding="utf-8"))["hata"]


def test_asil_ayri_derlemde_pilot_sorgularini_dislar(hazir, tmp_path):
    ozet = hazir / "ozet" / "egitim_havuz_ozeti_egitim_pilot.json"
    pilot_id = set(json.load(open(ozet, encoding="utf-8"))["pilot_sorgu_kaynak_kimlikleri"])
    assert len(pilot_id) == 5
    assert eh.ana(["--derlem", str(hazir / "derlem.jsonl"), "--onnx-klasor", str(hazir / "model"), "--asama", "egitim_asil",
                   "--cikti", str(tmp_path / "g"), "--ozet", str(tmp_path / "o"), "--n", "50", "--pilot-ozeti", str(ozet)]) == 0
    a = json.load(open(tmp_path / "g" / "yalnizca_arastirmaci" / "insan_anahtar.json", encoding="utf-8"))
    assert not {v["kaynak_kimligi"] for v in a["olusturma"]["sorgular"].values()} & pilot_id