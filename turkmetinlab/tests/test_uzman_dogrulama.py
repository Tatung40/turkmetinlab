"""Uzman doğrulama altyapısı: ölçütler (bilinen değerler), havuz körlemesi, dosya eşleştirme ve hata denetimleri.
Sentetik veri ve küçük sınama modeliyle çalışır; hiçbir gerçek uzman puanı içermez."""
import json
import os
import shutil
import subprocess
import sys

import numpy as np
import pytest
from openpyxl import load_workbook
from sklearn.metrics import cohen_kappa_score

from anlamsal_model import kucuk_model_olustur
from uzman_dogrulama import analiz
from uzman_dogrulama import ortak as o

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIZIBILITE = next((y for y in (os.path.join(KOK, "fizibilite_2.5"), os.path.join(os.path.dirname(KOK), "fizibilite_2.5"))
                   if os.path.isdir(y)), None)


# --- Ölçütler ------------------------------------------------------------------

@pytest.mark.parametrize("tur,sk", [("yok", None), ("dogrusal", "linear"), ("karesel", "quadratic")])
def test_kappa_scikit_learn_ile_ayni(tur, sk):
    rng = np.random.default_rng(3)
    a = rng.integers(0, 3, 300)
    b = np.where(rng.random(300) < 0.6, a, rng.integers(0, 3, 300))
    assert o.cohen_kappa(a, b, tur) == pytest.approx(cohen_kappa_score(a, b, weights=sk), abs=1e-12)


def _ac_bagimsiz(a, b, W):
    """Gwet (2014) formüllerinin döngüyle, ortak.py'den bağımsız yazımı."""
    n, q = len(a), 3
    pa = sum(W[x][y] for x, y in zip(a, b)) / n
    pi = [(list(a).count(k) + list(b).count(k)) / (2 * n) for k in range(q)]
    tw = sum(sum(r) for r in W)
    pe = tw / (q * (q - 1)) * sum(p * (1 - p) for p in pi)
    return (pa - pe) / (1 - pe)


def test_gwet_ac1_ac2_bilinen_degerler():
    rng = np.random.default_rng(5)
    a = rng.integers(0, 3, 200)
    b = np.where(rng.random(200) < 0.7, a, rng.integers(0, 3, 200))
    for tur in ("yok", "dogrusal", "karesel"):
        assert o.gwet_ac(a, b, tur) == pytest.approx(_ac_bagimsiz(a, b, o.agirliklar(tur).tolist()), abs=1e-12)
    assert o.gwet_ac(a, a) == pytest.approx(1) and o.cohen_kappa(a, a) == pytest.approx(1)
    # Marjinaller eşit dağılımlıysa AC1 = ağırlıksız κ (her ikisinde pe = 1/3)
    u = np.repeat([0, 1, 2], 10)
    v = u.copy()
    v[[0, 10, 20]] = [1, 2, 0]  # döngüsel kaydırma: marjinaller 10/10/10 kalır
    assert sorted(np.bincount(v)) == [10, 10, 10]
    assert o.gwet_ac(u, v, "yok") == pytest.approx(o.cohen_kappa(u, v, "yok"))


def test_ndcg_ve_precision_elle_hesaplanan():
    et = {"a": 2, "b": 0, "c": 1, "d": 2, "e": 0}
    dcg = 0 / np.log2(2) + 3 / np.log2(3) + 1 / np.log2(4)  # sıra: b, a, c
    idcg = 3 / np.log2(2) + 3 / np.log2(3) + 1 / np.log2(4)  # ideal: 2, 2, 1
    assert o.ndcg(["b", "a", "c"], et, 3) == pytest.approx(dcg / idcg)
    assert o.ndcg(["a", "d", "c"], et, 3) == pytest.approx(1)
    assert np.isnan(o.ndcg(["b"], {"b": 0}, 5))  # ilgili belge yok: tanımsız
    assert o.precision(["a", "b", "c", "d", "e"], et, 5, 2) == pytest.approx(0.4)
    assert o.precision(["a", "b", "c", "d", "e"], et, 5, 1) == pytest.approx(0.6)
    with pytest.raises(o.DogrulamaHatasi):  # yargısız belge ilgisiz sayılmaz
        o.ndcg(["a", "x"], et, 2)


def test_eslestirilmis_bootstrap():
    x = np.array([0.5, 0.6, np.nan, 0.7])
    r = o.eslestirilmis_bootstrap(x, x - 0.1)
    assert r["fark"] == pytest.approx(0.1) and r["ga"] == pytest.approx([0.1, 0.1]) and r["n_sorgu"] == 3
    rng = np.random.default_rng(0)
    y = rng.normal(0.05, 0.1, 50)
    r1, r2 = o.eslestirilmis_bootstrap(y, np.zeros(50)), o.eslestirilmis_bootstrap(y, np.zeros(50))
    assert r1 == r2 and r1["ga"][0] < r1["fark"] < r1["ga"][1]


# --- Uçtan uca: sentetik derlem → havuz → uzman dosyaları → analiz ----------------

KONU = [("güneş enerjisi panel verim", "fotovoltaik, güneş"), ("beton dayanım deney", "beton, yapı"),
        ("süt verimi hayvan besleme", "süt, hayvancılık"), ("deprem hasar bina", "deprem, yapı"),
        ("trafik akış kavşak", "trafik, ulaşım"), ("su arıtma membran", "su, arıtma")]


@pytest.fixture(scope="module")
def havuz(tmp_path_factory):
    if FIZIBILITE is None:
        pytest.skip("fizibilite_2.5 klasörü bulunamadı")
    d = tmp_path_factory.mktemp("uzman")
    kucuk_model_olustur(str(d / "model"), azami_token=4000)
    rng = np.random.default_rng(1)
    satir = []
    for i in range(700):
        konu, anahtar = KONU[i % len(KONU)]
        ek = " ".join(rng.choice(konu.split() + ["analiz", "yöntem", "sonuç", "çalışma"], 40))
        baslik = f"Makale {i} {konu}" if i != 7 else "Makale 1 güneş enerjisi panel verim"  # 7: 1'in başlık kopyası
        satir.append(f"{i}\t{baslik}\t{konu} üzerine çalışma {i}. {ek}\t{anahtar}, k{i % 3}")
    (d / "TRAll.txt").write_text("\n".join(satir), encoding="utf-8")
    for asama, klasor in (("asil", "gizli"), ("pilot", "gizli_pilot")):
        _havuz_calistir(d, asama, klasor)
    return d


def _havuz_calistir(d, asama, klasor, ozet="ozet"):
    r = subprocess.run([sys.executable, os.path.join(KOK, "uzman_dogrulama", "havuz_olustur.py"), "--onnx-klasor",
                        str(d / "model"), "--veri", str(d / "TRAll.txt"), "--fizibilite", FIZIBILITE, "--cikti",
                        str(d / klasor), "--ozet", str(d / ozet), "--n", "12", "--kume", "4", "--asama", asama],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stderr[-1500:]
    return r


def _anahtar(d):
    return json.load(open(d / "gizli" / "yalnizca_arastirmaci" / "anahtar.json", encoding="utf-8"))


def _uzman_yolu(d, u):
    return d / "gizli" / "uzmanlara_gonderin" / f"uzman_{u}.xlsx"


def test_havuz_tekillik_sizinti_ve_kapsam(havuz):
    a = _anahtar(havuz)
    ciftler = a["ciftler"]
    assert len({(c["sorgu_belge"], c["aday_belge"]) for c in ciftler}) == len(ciftler)  # mükerrer çift yok
    assert len({c["cift"] for c in ciftler}) == len(ciftler)
    assert all(c["sorgu_belge"] != c["aday_belge"] for c in ciftler)  # sorgu kendisini bulamaz
    for s, sira in a["siralamalar"].items():
        hepsi = set().union(*sira.values())
        assert hepsi == {c["cift"] for c in ciftler if str(c["sorgu_no"]) == s}  # havuz = ilk 10'ların birleşimi
        assert all(len(v) == o.ILK_K for v in sira.values())
    sorgu_belge = {c["sorgu_belge"] for c in ciftler}
    assert sorted({c["sorgu_no"] for c in ciftler}) == list(range(1, 13))
    sys.path.append(FIZIBILITE)
    import anlamsal_karsilastirma as ak
    _, ozet, _, n = ak.veri_oku(str(havuz / "TRAll.txt"), 2)
    tum = ak.veri_oku(str(havuz / "TRAll.txt"), n)[1]
    onceki = set(ak.veri_oku(str(havuz / "TRAll.txt"), 500)[1])
    assert not {tum[s] for s in sorgu_belge} & onceki  # önceki örneklemden sorgu yok
    if 1 in sorgu_belge:  # başlık kopyası (7) adaylarda olmamalı
        assert 7 not in {c["aday_belge"] for c in ciftler if c["sorgu_belge"] == 1}


def test_uzman_dosyalari_kor_ve_kilitli(havuz):
    a = _anahtar(havuz)
    sira = {}
    for u in (1, 2):
        wb = load_workbook(_uzman_yolu(havuz, u))
        assert wb.sheetnames == [o.YONERGE_SAYFASI, o.SAYFA] and all(ws.sheet_state == "visible" for ws in wb)
        assert not wb.defined_names
        ws = wb[o.SAYFA]
        satirlar = list(ws.iter_rows(values_only=True))
        assert list(satirlar[0]) == o.SUTUNLAR and ws.protection.sheet
        assert ws.cell(2, 1).protection.locked and not ws.cell(2, 7).protection.locked
        assert any("0,1,2" in dv.formula1 for dv in ws.data_validations.dataValidation)
        for ws_ in wb:  # yöntem adı, sıra ya da puan yok (sentetik metinler bu sözcükleri içermez)
            for r in ws_.iter_rows(values_only=True):
                for h in r:
                    assert not any(y in str(h) for y in o.YASAK_IFADELER if h is not None), h
                    assert not isinstance(h, float)
        assert {r[0] for r in satirlar[1:]} == {c["cift"] for c in a["ciftler"]}
        sira[u] = [r[0] for r in satirlar[1:]]
        assert [r[1] for r in satirlar[1:]] == sorted(r[1] for r in satirlar[1:])  # sorgu bloklar halinde
    assert sira[1] != sira[2]  # uzmanlar için ayrı karıştırma


def _doldur(havuz, u, puan_fn, hedef, klasor="gizli", on=""):
    wb = load_workbook(havuz / klasor / "uzmanlara_gonderin" / f"{on}uzman_{u}.xlsx")
    ws = wb[o.SAYFA]
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 7).value = puan_fn(ws.cell(r, 1).value)
    wb.save(hedef)
    return hedef


def test_analiz_bilinen_sonuc_ve_uzlasma(havuz, tmp_path):
    a = _anahtar(havuz)
    hib = {c for s in a["siralamalar"].values() for c in s["Hibrit"][:3]}
    tf = {c for s in a["siralamalar"].values() for c in s["TF-IDF"]}
    gercek = {c["cift"]: (2 if c["cift"] in hib and c["cift"] not in tf else 0) for c in a["ciftler"]}
    p1 = _doldur(havuz, 1, gercek.get, tmp_path / "u1.xlsx")
    ilk = sorted(gercek)[:5]  # uzman 2: beş çiftte farklı
    p2 = _doldur(havuz, 2, lambda c: (1 if c in ilk else gercek[c]), tmp_path / "u2.xlsx")
    anahtar = str(havuz / "gizli" / "yalnizca_arastirmaci" / "anahtar.json")
    assert analiz.ana(["--anahtar", anahtar, "--uzman1", str(p1), "--uzman2", str(p2), "--cikti", str(tmp_path / "s")]) == 0
    r = json.load(open(tmp_path / "s" / "analiz.json", encoding="utf-8"))
    assert r["uyum"]["uyusmazlik"]["toplam"] == 5 and "ön analiz" in r["asama"]
    d = r["uyum"]["dagilim"]["uzman_1"]
    assert sum(v["n"] for v in d.values()) == len(a["ciftler"]) and sum(v["yuzde"] for v in d.values()) == pytest.approx(100, abs=0.05)
    e = r["erisim"]["uzman_1"]
    assert e["ortalama"]["TF-IDF"]["p@10_kati"] == 0  # TF-IDF'nin bulduğu hiçbir çift ilgili değil
    assert e["farklar (eşleştirilmiş bootstrap, %95 GA)"]["Hibrit − TF-IDF"]["p@10_kati"]["fark"] >= 0
    assert "Karar verilmez" in r["karar"]["durum"]
    wb = load_workbook(tmp_path / "s" / "uzlasma.xlsx")
    ws = wb["Uzlaşma"]
    assert ws.max_row - 1 == 5
    for i in range(2, ws.max_row + 1):
        ws.cell(i, 11).value = gercek[ws.cell(i, 1).value]
    wb.save(tmp_path / "uz.xlsx")
    assert analiz.ana(["--anahtar", anahtar, "--uzman1", str(p1), "--uzman2", str(p2), "--uzlasma",
                       str(tmp_path / "uz.xlsx"), "--cikti", str(tmp_path / "s2")]) == 0
    r2 = json.load(open(tmp_path / "s2" / "analiz.json", encoding="utf-8"))
    assert r2["asama"].startswith("birincil") and "pratik_onem" in r2["karar"]
    assert r2["erisim"]["uzlasi"]["ortalama"] == r["erisim"]["uzman_1"]["ortalama"]  # uzlaşı = gerçek etiket


@pytest.mark.parametrize("bozma,ileti", [
    ("ayni_dosya", "Uzman kodları"), ("eksik_puan", "Eksik"), ("gecersiz_puan", "geçersiz"),
    ("metin_degisti", "denetim"), ("yinelenen", "yinelenen"), ("baska_havuz", "havuz")])
def test_eslestirme_ve_girdi_hatalari(havuz, tmp_path, bozma, ileti):
    anahtar = str(havuz / "gizli" / "yalnizca_arastirmaci" / "anahtar.json")
    p1 = _doldur(havuz, 1, lambda c: 1, tmp_path / "u1.xlsx")
    p2 = _doldur(havuz, 2, lambda c: 1, tmp_path / "u2.xlsx")
    wb = load_workbook(p2)
    ws, ys = wb[o.SAYFA], wb[o.YONERGE_SAYFASI]
    if bozma == "ayni_dosya":
        p2 = p1
    elif bozma == "eksik_puan":
        ws.cell(5, 7).value = None
    elif bozma == "gecersiz_puan":
        ws.cell(5, 7).value = 3
    elif bozma == "metin_degisti":
        ws.cell(5, 6).value = str(ws.cell(5, 6).value) + " ek"
    elif bozma == "yinelenen":
        ws.cell(5, 1).value = ws.cell(6, 1).value
    elif bozma == "baska_havuz":
        ys["A2"].value = "Havuz kimliği: 0000000000000000"
    if bozma != "ayni_dosya":
        wb.save(p2)
    assert analiz.ana(["--anahtar", anahtar, "--uzman1", str(p1), "--uzman2", str(p2), "--cikti", str(tmp_path / "s")]) == 2
    assert ileti.lower() in json.load(open(tmp_path / "s" / "durum.json", encoding="utf-8"))["hata"].lower()


# --- Pilot ---------------------------------------------------------------------

def _pilot_anahtar(d):
    return json.load(open(d / "gizli_pilot" / "yalnizca_arastirmaci" / "pilot_anahtar.json", encoding="utf-8"))


def test_pilot_ayri_ve_asil_ile_ortusmez(havuz):
    p, a = _pilot_anahtar(havuz), _anahtar(havuz)
    assert p["asama"] == "pilot" and a["asama"] == "asil" and p["havuz_kimligi"] != a["havuz_kimligi"]
    assert len({c["sorgu_no"] for c in p["ciftler"]}) == o.PILOT_SAYISI
    assert all(c["cift"].startswith("P-") for c in p["ciftler"]) and all(c["cift"].startswith("Ç-") for c in a["ciftler"])
    assert not {c["sorgu_belge"] for c in p["ciftler"]} & {c["sorgu_belge"] for c in a["ciftler"]}
    assert sorted(os.listdir(havuz / "gizli_pilot" / "uzmanlara_gonderin")) == ["pilot_uzman_1.xlsx", "pilot_uzman_2.xlsx"]
    ys = load_workbook(havuz / "gizli_pilot" / "uzmanlara_gonderin" / "pilot_uzman_1.xlsx")[o.YONERGE_SAYFASI]
    assert "PİLOT" in ys["A3"].value
    ozet = json.load(open(havuz / "ozet" / "uzman_havuz_ozeti_pilot.json", encoding="utf-8"))
    assert "BİLİNMİYOR" in ozet["kaynak_lisans_durumu"] and ozet["pilot_asil_sorgu_ortusmesi"] == 0
    assert not [f for f in os.listdir(havuz / "ozet") if not f.startswith("uzman_havuz_ozeti_")]  # özette metin/anahtar yok


def test_asil_havuz_pilottan_etkilenmez(havuz, tmp_path):
    shutil.copytree(havuz / "model", tmp_path / "model")
    shutil.copy(havuz / "TRAll.txt", tmp_path / "TRAll.txt")
    _havuz_calistir(tmp_path, "asil", "g", "oz")
    assert json.load(open(tmp_path / "g" / "yalnizca_arastirmaci" / "anahtar.json"))["ciftler"] == _anahtar(havuz)["ciftler"]


def test_pilot_analizi_karar_vermez_ve_karistirilamaz(havuz, tmp_path):
    p = _pilot_anahtar(havuz)
    pa = str(havuz / "gizli_pilot" / "yalnizca_arastirmaci" / "pilot_anahtar.json")
    u = {i: _doldur(havuz, i, lambda c: 1, tmp_path / f"p{i}.xlsx", "gizli_pilot", "pilot_") for i in (1, 2)}
    assert analiz.ana(["--anahtar", pa, "--uzman1", str(u[1]), "--uzman2", str(u[2]), "--cikti", str(tmp_path / "s")]) == 0
    r = json.load(open(tmp_path / "s" / "pilot_analiz.json", encoding="utf-8"))
    assert r["calisma"] == "pilot" and "PİLOT" in r["karar"]["durum"] and not (tmp_path / "s" / "analiz.json").exists()
    assert len(r["erisim"]["uzman_1"]["ortalama"]) == 3 and r["uyum"]["cift_sayisi"] == len(p["ciftler"])
    asil = str(havuz / "gizli" / "yalnizca_arastirmaci" / "anahtar.json")  # pilot dosyası asıl anahtarla analiz edilemez
    assert analiz.ana(["--anahtar", asil, "--uzman1", str(u[1]), "--uzman2", str(u[2]), "--cikti", str(tmp_path / "x")]) == 2
    assert "aşama" in json.load(open(tmp_path / "x" / "durum.json", encoding="utf-8"))["hata"]


def test_gizli_ve_acik_klasor_ayni_olamaz(havuz, tmp_path):
    shutil.copytree(havuz / "model", tmp_path / "model")
    shutil.copy(havuz / "TRAll.txt", tmp_path / "TRAll.txt")
    with pytest.raises(AssertionError):
        _havuz_calistir(tmp_path, "asil", "ayni", "ayni")


@pytest.mark.skipif(shutil.which("7z") is None, reason="7z yok (GitHub Actions'ta kurulur)")
def test_paketleme_sifreli_ve_dogrulanmis(havuz, tmp_path):
    hedef = tmp_path / "p.7z"
    betik = os.path.join(KOK, "uzman_dogrulama", "paketle.sh")
    r = subprocess.run(["bash", betik, str(havuz / "gizli"), str(hedef)], capture_output=True, text=True,
                       env={**os.environ, "PAROLA": "deneme-parolasi-123"})
    assert r.returncode == 0 and hedef.exists() and "deneme-parolasi" not in r.stdout + r.stderr
    assert subprocess.run(["7z", "l", "-pyanlis", str(hedef)], capture_output=True).returncode != 0
    r = subprocess.run(["bash", betik, str(havuz / "gizli"), str(tmp_path / "q.7z")], capture_output=True, text=True,
                       env={**os.environ, "PAROLA": ""})
    assert r.returncode != 0 and not (tmp_path / "q.7z").exists()