"""DOCX belge okuma ve güvenlik regresyon testleri."""
import io
import pytest
from docx import Document
from tmlab.loader import VeriOkumaHatasi, docx_oku, dosya_oku


def _word(paragraflar=(), tablo=False):
    b=io.BytesIO()
    d=Document()
    for p in paragraflar:
        d.add_paragraph(p)
    if tablo:
        t=d.add_table(rows=2, cols=2)
        t.cell(0,0).text="Katılımcı"
        t.cell(0,1).text="Görüş"
        t.cell(1,0).text="K01"
        t.cell(1,1).text="Uzaktan eğitim yararlı."
    d.save(b)
    return b.getvalue()


def test_docx_turkce_ve_tablo():
    s=docx_oku([("anket.docx", _word(["Işık ve ölçme.", "Öğrencilerin görüşleri."], True))])
    assert len(s.tablo)==1
    assert "Işık ve ölçme." in s.tablo.iloc[0]["metin"]
    assert "K01 | Uzaktan eğitim yararlı." in s.tablo.iloc[0]["metin"]
    assert s.tablo.iloc[0]["belge"]=="anket.docx"


def test_docx_paragraf_bolme_ve_tablo_satirlari():
    s=docx_oku([("a.docx", _word(["Birinci.", "İkinci."], True))], "paragraf")
    assert len(s.tablo)==4
    assert s.tablo.iloc[2]["metin"]=="Katılımcı | Görüş"


def test_docx_coklu_dosya():
    s=docx_oku([("a.docx",_word(["A"])), ("b.docx",_word(["B"]))])
    assert s.tablo["belge"].tolist()==["a.docx","b.docx"]


def test_docx_uzanti_ile_oku():
    assert dosya_oku("deneme.docx",_word(["Test"])).tablo.iloc[0]["metin"]=="Test"


@pytest.mark.parametrize("veri",[b"",b"metin",b"PK\\x03\\x04broken"])
def test_docx_bozuk_reddedilir(veri):
    with pytest.raises(VeriOkumaHatasi):
        docx_oku([("bozuk.docx",veri)])


def test_docx_yanlis_mod_reddedilir():
    with pytest.raises(VeriOkumaHatasi):
        docx_oku([("x.docx",_word(["test"]))],"hatalı")


def test_docx_bos_reddedilir():
    with pytest.raises(VeriOkumaHatasi):
        docx_oku([("bos.docx",_word())])
