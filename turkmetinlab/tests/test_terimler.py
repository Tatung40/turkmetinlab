"""2.1 — Akademik terim sözlüğü ve çok sözcüklü terim tanıma testleri.

Bu dosyadaki örnekler GELİŞTİRME verisidir (entegrasyon notlarındaki örnekler ve elle yazılmış dilbilgisi
denetimleri). Bağımsız değerlendirme için ayrı tutulan yapay test kümesi degerlendirme/veri/terim_test_yapay.tsv
dosyasındadır ve test_terim_test_kumesi_* testlerinde yalnızca bir eşik altına düşülmediği denetlenir.
"""

import io
import os
import re

import pandas as pd
import pytest

from tmlab import terimler as tm
from tmlab.analysis import ngram_sikligi, sozcuk_sikligi, tfidf_hesapla
from tmlab.birliktelik import AgAyarlari, birliktelik_hesapla
from tmlab.preprocess import OnIslemeAyarlari, derlem_on_isle, on_isle
from tmlab.terimler import (TerimEslestirici, Terim, cekirdek_sozluk, kullanici_terimi, son_sozcuk_cozumle,
                            sozluk_olustur, sozluk_tablosundan)

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


@pytest.fixture(scope="module")
def E():
    return TerimEslestirici(cekirdek_sozluk())


@pytest.fixture(scope="module")
def EK():
    return TerimEslestirici(cekirdek_sozluk(), kisaltmalar=True)


def bul(E, metin):
    r = E.eslestir(metin)
    for e in r.eslesmeler:  # konumlar her zaman özgün metne geri eşlenir
        assert metin[e.bas:e.bit] == e.yuzey
    return [(e.yuzey, e.standart) for e in r.eslesmeler]


# --- Sözlük -----------------------------------------------------------------

def test_cekirdek_sozluk_yapisi():
    s = cekirdek_sozluk()
    assert len(s) == 637 and len({t.terim_id for t in s}) == 637
    assert len(tm.alanlar()) == 6 and tm.alanlar()[-1] == tm.GENEL_ALAN
    assert all(t.durum.startswith("aday") for t in s)
    assert sum(bool(t.varyantlar) for t in s) == 20
    assert "637 aday terim" in tm.sozluk_surumu() and "SHA-256" in tm.sozluk_surumu()


def test_alan_secimi_ve_devre_disi():
    egitim = "Eğitim bilimleri ve psikometri"
    s = sozluk_olustur([egitim])
    assert len(s) == 114 and {t.alan for t in s} == {egitim}
    s2 = sozluk_olustur([egitim], kapali=["T0001"])
    assert len(s2) == 113 and "T0001" not in {t.terim_id for t in s2}
    E = TerimEslestirici(s2)
    assert bul(E, "uzaktan eğitimde") == []
    # Seçilmeyen alanın terimi tanınmaz
    assert bul(TerimEslestirici(s), "yapay zekâ") == []


# --- Entegrasyon notlarındaki ve istemdeki örnekler ------------------------------

@pytest.mark.parametrize("metin,beklenen", [
    ("Uzaktan eğitime katılan öğrenci", [("Uzaktan eğitime", "uzaktan eğitim")]),
    ("Yapay zekânın etkileri", [("Yapay zekânın", "yapay zekâ")]),
    ("Ölçme ve değerlendirmede yenilik", [("Ölçme ve değerlendirmede", "ölçme ve değerlendirme")]),
    ("Yapısal eşitlik modellemesinin sonuçları",
     [("Yapısal eşitlik modellemesinin", "yapısal eşitlik modellemesi")]),
    ("madde tepki kuramına göre", [("madde tepki kuramına", "madde tepki kuramı")]),
    ("Uzaktan eğitimci bir araştırmacı", []),
    ("Yapay zekâsız bir ortam", []),
    ("ölçme ve değerlendirmeci", []),
])
def test_notlardaki_ornekler(E, metin, beklenen):
    assert bul(E, metin) == beklenen


def test_turetim_bicimleri_ayrica_kaydedilir(E):
    r = E.eslestir("Uzaktan eğitimci ve yapay zekâsız ortamlar")
    nedenler = {(x.yuzey, x.neden) for x in r.eslesmeyenler}
    assert ("Uzaktan eğitimci", "türetim eki") in nedenler and ("yapay zekâsız", "türetim eki") in nedenler


@pytest.mark.parametrize("ifade,standart", [
    ("uzaktan eğitim", "uzaktan eğitim"), ("yapay zekâ", "yapay zekâ"), ("yapay zeka", "yapay zekâ"),
    ("ölçme ve değerlendirme", "ölçme ve değerlendirme"), ("ölçme-değerlendirme", "ölçme ve değerlendirme"),
    ("yapısal eşitlik modellemesi", "yapısal eşitlik modellemesi"), ("madde tepki kuramı", "madde tepki kuramı"),
    ("doğal dil işleme", "doğal dil işleme")])
@pytest.mark.parametrize("ek_bicim", ["{}", "{}nin", "{}de", "{}den", "{}ye", "{}yi", "{}yle", "{}deki", "{}dir"])
def test_zorunlu_terimler_cekimli_ve_cekimsiz(E, ifade, standart, ek_bicim):
    """Entegrasyon notlarının 8. maddesi: altı terim, ekli ve eksiz (ekler elle yazılmış doğru biçimler)."""
    dogru = {
        "uzaktan eğitim": ["uzaktan eğitim", "uzaktan eğitimin", "uzaktan eğitimde", "uzaktan eğitimden",
                           "uzaktan eğitime", "uzaktan eğitimi", "uzaktan eğitimle", "uzaktan eğitimdeki",
                           "uzaktan eğitimdir"],
        "yapay zekâ": ["yapay zekâ", "yapay zekânın", "yapay zekâda", "yapay zekâdan", "yapay zekâya", "yapay zekâyı",
                       "yapay zekâyla", "yapay zekâdaki", "yapay zekâdır"],
        "yapay zeka": ["yapay zeka", "yapay zekanın", "yapay zekada", "yapay zekadan", "yapay zekaya", "yapay zekayı",
                       "yapay zekayla", "yapay zekadaki", "yapay zekadır"],
        "ölçme ve değerlendirme": ["ölçme ve değerlendirme", "ölçme ve değerlendirmenin", "ölçme ve değerlendirmede",
                                   "ölçme ve değerlendirmeden", "ölçme ve değerlendirmeye",
                                   "ölçme ve değerlendirmeyi", "ölçme ve değerlendirmeyle",
                                   "ölçme ve değerlendirmedeki", "ölçme ve değerlendirmedir"],
        "ölçme-değerlendirme": ["ölçme-değerlendirme", "ölçme-değerlendirmenin", "ölçme-değerlendirmede",
                                "ölçme-değerlendirmeden", "ölçme-değerlendirmeye", "ölçme-değerlendirmeyi",
                                "ölçme-değerlendirmeyle", "ölçme-değerlendirmedeki", "ölçme-değerlendirmedir"],
        "yapısal eşitlik modellemesi": ["yapısal eşitlik modellemesi", "yapısal eşitlik modellemesinin",
                                        "yapısal eşitlik modellemesinde", "yapısal eşitlik modellemesinden",
                                        "yapısal eşitlik modellemesine", "yapısal eşitlik modellemesini",
                                        "yapısal eşitlik modellemesiyle", "yapısal eşitlik modellemesindeki",
                                        "yapısal eşitlik modellemesidir"],
        "madde tepki kuramı": ["madde tepki kuramı", "madde tepki kuramının", "madde tepki kuramında",
                               "madde tepki kuramından", "madde tepki kuramına", "madde tepki kuramını",
                               "madde tepki kuramıyla", "madde tepki kuramındaki", "madde tepki kuramıdır"],
        "doğal dil işleme": ["doğal dil işleme", "doğal dil işlemenin", "doğal dil işlemede", "doğal dil işlemeden",
                             "doğal dil işlemeye", "doğal dil işlemeyi", "doğal dil işlemeyle", "doğal dil işlemedeki",
                             "doğal dil işlemedir"],
    }[ifade]
    sira = ["{}", "{}nin", "{}de", "{}den", "{}ye", "{}yi", "{}yle", "{}deki", "{}dir"].index(ek_bicim)
    metin = f"Bu çalışmada {dogru[sira]} ele alındı."
    assert bul(E, metin) == [(dogru[sira], standart)]


# --- Dilbilgisi: kabul edilen ve reddedilen biçimler (elle yazılmış listeler) ---

@pytest.mark.parametrize("w,n,kabul,ret", [
    ("eğitim", 2,
     "eğitime eğitimde eğitimden eğitimin eğitimi eğitimle eğitimler eğitimlerin eğitimlerde eğitimimiz "
     "eğitiminde eğitimdeki eğitimdir eğitimdi eğitimken eğitimlerimizdeki eğitimindeki eğitimdekiler",
     "eğitimci eğitimli eğitimsel eğitimsiz eğitimlik eğitme eğitimcilik eğitimda eğitimlar eğitimsizlik"),
    ("kuramı", 3,
     "kuramına kuramının kuramından kuramında kuramını kuramıyla kuramındaki kuramları kuramlarının "
     "kuramlarında kuramıdır",
     # Not: "kuramıda" gibi Türkçede geçmeyen biçimler bilerek reddedilmez: "tomografide" gibi -i ile biten
     # sözlük biçimleri için normal hâl ekleri de gerekir; geçmeyen bir biçimi kabul etmek yanlış pozitif üretmez.
     "kuramcı kuramsal kuramsallaştırma kuramsız kuramine kuramına_"),
    ("zeka", 2, "zekanın zekaya zekada zekası zekasının zekalar zekayla zekadır",
     "zekasız zekalı zekice zekanin zekade"),
    ("sağlığı", 2, "sağlığının sağlığına sağlığında sağlıkları sağlıklarının sağlığıyla",
     "sağlıklı sağlıksız sağlıkçı sağlıklılık"),
    ("kaydı", 3, "kaydının kayıtları kayıtlarında kaydına", "kayıtlı kayıtsız"),
    ("kontrolü", 2, "kontrolünün kontrolleri kontrollerinin kontrolünde", "kontrolsüz kontrollü kontrolcü"),
    ("tıp", 2, "tıbbın tıbba tıbbı tıpta tıptaki", "tıbbi tıpçı"),
    ("analiz", 2, "analizi analizler analizlerin analizde analizinde", "analizci analizsel"),
    ("risk", 1, "riskin riske riskler riskte", "riskli risksiz riskçi"),
])
def test_cekim_ve_turetim_ayrimi(w, n, kabul, ret):
    for t in kabul.split():
        assert son_sozcuk_cozumle(t, None, w, n)[0] is not None, t
    for t in ret.split():
        assert son_sozcuk_cozumle(t, None, w, n)[0] is None, t


def test_turetim_nedenleri():
    assert son_sozcuk_cozumle("eğitimci", None, "eğitim", 2) == (None, "türetim eki")
    assert son_sozcuk_cozumle("eğitimsel", None, "eğitim", 2) == (None, "türetim eki")
    assert son_sozcuk_cozumle("eğitimxq", None, "eğitim", 2) == (None, "tanınmayan ek")
    assert son_sozcuk_cozumle("öğretim", None, "eğitim", 2) == (None, "")  # ilgisiz sözcük


def test_sozcugun_kendisi_i_ile_biten_terimler(E):
    assert bul(E, "bilgisayarlı tomografide ve bilgisayarlı tomografinin") == [
        ("bilgisayarlı tomografide", "bilgisayarlı tomografi"), ("bilgisayarlı tomografinin", "bilgisayarlı tomografi")]


def test_genel_unlu_dusmesi_yanlis_pozitif_uretmez(E):
    # "eğitim → eğitme", "ölçüm → ölçme" gibi gerçek ama farklı sözcükler eşleşmemeli
    assert bul(E, "uzaktan eğitme yöntemleri") == []
    assert bul(E, "hizmet içi eğitme") == []


def test_kesme_isareti(E):
    assert bul(E, "Kaplan Meier'in yöntemi") == []  # terim tamamlanmadı
    assert bul(E, "Yapay zeka'nın ve yapay zekâ’ya") == [("Yapay zeka'nın", "yapay zekâ"),
                                                         ("yapay zekâ’ya", "yapay zekâ")]
    assert bul(E, "yapay zeka'cı") == []


# --- Sözcük sınırları ve noktalama -------------------------------------------

@pytest.mark.parametrize("metin", [
    "Model, açıklanabilir sonuçlar üretti.", "risk; faktör", "Sinir, ağır yük taşıyan işçiler.",
    "uzaktan\n\neğitim", "ölçme aracı ve değerlendirme ölçütleri", "Uzak eğitim merkezleri",
    "Bu ayar ince bir dengeyi gerektiriyor.", "derinöğrenme", "yapayzekâ",
    "Ayrıntılar https://ornek.org/yapay-zeka/ adresinde.",
])
def test_yanlis_pozitif_yok(E, metin):
    assert bul(E, metin) == []


def test_tireli_bilesigin_parcasi_eslesmez(E):
    r = E.eslestir("log-normal dağılım modelleri ve normal dağılım")
    assert [e.yuzey for e in r.eslesmeler] == ["normal dağılım"] and r.eslesmeler[0].bas > 20
    assert any(x.neden == "tireli bileşiğin parçası" for x in r.eslesmeyenler)


def test_iyelik_dusmesi_zincirleme_tamlama(E):
    assert bul(E, "klasik makine öğrenme yöntemleri") == [("makine öğrenme", "makine öğrenmesi")]
    assert bul(E, "Pearson korelasyon katsayısı") == [("Pearson korelasyon", "Pearson korelasyonu")]
    r = E.eslestir("saldırı tespit sistemleri")
    assert r.eslesmeler[0].ek == "iyeliksiz" and r.eslesmeler[0].guven == 0.75
    # Ardından sözcük gelmiyorsa (bölüt sonu) kabul edilmez
    assert bul(E, "Makine Öğrenme: Rastgele orman") == [("Rastgele orman", "rastgele orman")]
    assert bul(E, "makine öğrenme.") == []


def test_kesme_oncesi_cekim_tek_harfli_onek_ve_iyeliksiz_ekli(E):
    # Gerçek veride (DergiPark) görülen biçimler
    assert bul(E, "Evrişimli Sinir Ağları’dır.") == [("Evrişimli Sinir Ağları’dır", "evrişimli sinir ağı")]
    assert bul(E, "k-En yakın komşu algoritması") == [("En yakın komşu algoritması", "en yakın komşu algoritması")]
    r = E.eslestir("yaşam kalitemizi ve makine öğrenmenin katkısı")
    assert [(e.standart, e.ek) for e in r.eslesmeler] == [("yaşam kalitesi", "iyeliksiz"), ("makine öğrenmesi", "iyeliksiz")]
    assert bul(E, "anne sağlıklı bebekler") == []  # türetim eki iyeliksiz gövdede de kabul edilmez


def test_tire_satir_sonu_ve_cok_bosluk(E):
    assert bul(E, "meta-analizin bulguları") == [("meta-analizin", "meta analiz")]
    assert bul(E, "Kaplan-Meier eğrisi") == [("Kaplan-Meier eğrisi", "Kaplan Meier eğrisi")]
    assert bul(E, "uzaktan\neğitimde") == [("uzaktan\neğitimde", "uzaktan eğitim")]
    assert E.eslestir("uzaktan\neğitimde", satir_sonu_sinir=True).eslesmeler == []
    assert bul(E, "uzaktan   eğitimde") == [("uzaktan   eğitimde", "uzaktan eğitim")]
    assert bul(E, "uzaktan - eğitim") == []  # boşluklu tire bölüt sınırıdır


def test_konumlar_unicode_normallestirmede_korunur(E):
    # Ayrık yazılmış aksan (e + U+0301 değil; ğ = g + U+0306), yumuşak tire, bitişik harf (ﬁ), satır sonu tiresi
    metinler = ["Uzaktan eğitime katılım", "uzaktan eği­timde", "ﬁnansal istikrarın korunması",
                "UZAKTAN EĞİTİMİN", "uzaktan eği-\ntimin yaygınlaşması", "İŞSİZLİK ORANI"]
    beklenen = ["uzaktan eğitim", "uzaktan eğitim", "finansal istikrar", "uzaktan eğitim", "uzaktan eğitim",
                "işsizlik oranı"]
    for m, b in zip(metinler, beklenen):
        r = E.eslestir(m)
        assert [e.standart for e in r.eslesmeler] == [b], m
        e = r.eslesmeler[0]
        assert m[e.bas:e.bit] == e.yuzey and e.bas == 0


# --- İç içe ve kesişen terimler, öncelik ------------------------------------------

def test_ic_ice_en_uzun_secilir(E):
    assert bul(E, "Üretken yapay zekâ araçları") == [("Üretken yapay zekâ", "üretken yapay zekâ")]
    assert bul(E, "algılanan sosyal desteğin") == [("algılanan sosyal desteğin", "algılanan sosyal destek")]
    assert bul(E, "rastgele eksik veri") == [("rastgele eksik veri", "rastgele eksik veri")]
    r = E.eslestir("yapay zekâ okuryazarlığının")
    assert [e.standart for e in r.eslesmeler] == ["yapay zekâ okuryazarlığı"]
    assert ("yapay zekâ", "iç içe") in {(c.reddedilen_standart, c.tur) for c in r.cakismalar}


def test_kesisen_terimler_kaydedilir(E):
    r = E.eslestir("sosyal girişimcilik ekosistemi")
    assert len(r.eslesmeler) == 1
    assert {c.tur for c in r.cakismalar} == {"kesişen"}
    alinan = r.eslesmeler[0].standart
    reddedilen = r.cakismalar[0].reddedilen_standart
    assert {alinan, reddedilen} == {"sosyal girişimcilik", "girişimcilik ekosistemi"}


def test_oncelik_kullanici_alan_genel():
    kendi = kullanici_terimi("uzaktan eğitim", terim_id="K1")
    E = TerimEslestirici(sozluk_olustur(kullanici=[kendi]))
    r = E.eslestir("uzaktan eğitimde")
    assert r.eslesmeler[0].terim_id == "K1" and r.eslesmeler[0].kaynak == "kullanıcı"
    assert r.cakismalar[0].tur == "aynı konum" and r.cakismalar[0].reddedilen_id == "T0001"
    # Aynı yüzeyde alan sözlüğü, Genel araştırma ve istatistik listesine üstün gelir
    a = Terim("A1", "Sağlık bilimleri", "örnek terim")
    g = Terim("G1", tm.GENEL_ALAN, "örnek terim")
    r = TerimEslestirici([g, a]).eslestir("örnek terimi")
    assert r.eslesmeler[0].terim_id == "A1"


def test_standart_varyanttan_once():
    t1 = Terim("A", "x", "çevrim içi öğrenme", ("çevrimiçi öğrenme",))
    r = TerimEslestirici([t1]).eslestir("çevrimiçi öğrenmede ve çevrim içi öğrenme")
    assert [(e.bicim, e.ek) for e in r.eslesmeler] == [("varyant", "çekim"), ("standart", "yok")]
    assert [e.guven for e in r.eslesmeler] == [0.85, 1.0]


# --- Kısaltmalar -----------------------------------------------------------

def test_kisaltmalar_varsayilan_kapali(E):
    assert bul(E, "Model YEM ile sınandı.") == []


def test_kisaltmalar_buyuk_kucuk_harf_duyarli(EK):
    assert bul(EK, "Model YEM ile sınandı.") == [("YEM", "yapısal eşitlik modellemesi")]
    assert bul(EK, "KEKK-YEM yöntemiyle") == [("YEM", "yapısal eşitlik modellemesi")]
    assert bul(EK, "YEM'in varsayımları") == [("YEM'in", "yapısal eşitlik modellemesi")]
    assert bul(EK, "hayvansal yem ve yem bitkisi") == []
    assert bul(EK, "Yem bitkileri") == []
    assert bul(EK, "Sensörler IoT ağına bağlandı.") == [("IoT", "nesnelerin interneti")]
    assert bul(EK, "YEMİ") == []
    assert bul(EK, "Bulanık Mantık ve IOT Kullanarak") == [("IOT", "nesnelerin interneti")]
    assert bul(EK, "iot cihazları") == []


def test_kisaltma_basliklarda_reddedilir(EK):
    r = EK.eslestir("HAYVAN YEM SANAYİSİ RAPORU\nBu raporda YEM kullanılmadı.")
    assert len(r.eslesmeler) == 1 and r.eslesmeler[0].bas > 20
    assert any(x.neden == "büyük harfli başlık satırı" for x in r.eslesmeyenler)


# --- Kullanıcı sözlüğü ---------------------------------------------------------

def test_kullanici_terimi_dogrulama():
    t = kullanici_terimi("  makine   öğrenimi ", "makine öğrenmesi; ML")
    assert t.standart == "makine öğrenimi" and t.varyantlar == ("makine öğrenmesi", "ML") and t.kaynak == "kullanıcı"
    for kotu in ["", "   ", "123 456", "a " * 12, "x" * 130]:
        with pytest.raises(tm.SozlukHatasi):
            kullanici_terimi(kotu)


def test_kullanici_terimi_eslesir_ve_cekim_alir():
    E = TerimEslestirici([kullanici_terimi("makine öğrenimi", terim_id="K1")])
    assert bul(E, "Makine öğreniminin ve makine öğrenimci") == [("Makine öğreniminin", "makine öğrenimi")]


def test_sozluk_tablosundan_yukleme():
    tablo = pd.DataFrame({"Standart_Terim": ["makine öğrenimi", "Makine Öğrenimi", None, "123", "işbirlikli öğrenme",
                                             "=HYPERLINK(\"x\")"],
                          "yazim_varyantlari": ["makine öğrenmesi", None, None, None, "iş birlikli öğrenme", None],
                          "alan": [None, None, None, None, "Eğitim", None],
                          "terim_id": ["T0001", None, None, None, "E7", None]})
    terimler, uyarilar = sozluk_tablosundan(tablo)
    assert [t.standart for t in terimler] == ["makine öğrenimi", "işbirlikli öğrenme", "=HYPERLINK(\"x\")"]
    assert terimler[0].terim_id == "Y0001"  # çekirdek kimliğiyle çakışan kimlik yeniden verilir
    assert terimler[1].terim_id == "E7" and terimler[1].alan == "Eğitim"
    assert any("birden fazla" in u for u in uyarilar) and any("geçersiz" in u for u in uyarilar)
    E = TerimEslestirici(terimler)
    assert bul(E, "işbirlikli öğrenmenin") == [("işbirlikli öğrenmenin", "işbirlikli öğrenme")]
    with pytest.raises(tm.SozlukHatasi):
        sozluk_tablosundan(pd.DataFrame({"standart_terim": [None, "  "]}))


def test_sozluk_dosyasi_guvenli_okuyucuyla_ve_disa_aktarim_formulsuz():
    from openpyxl import load_workbook

    from tmlab.export import excel_bayt
    from tmlab.loader import dosya_oku

    csv = "standart_terim;yazim_varyantlari\nmakine öğrenimi;makine öğrenmesi\n=TOPLA(A1);\n".encode("utf-8")
    okuma = dosya_oku("sozluk.csv", csv)
    terimler, _ = sozluk_tablosundan(okuma.tablo)
    tablo = tm.sozluk_disa_aktar(terimler)
    wb = load_workbook(io.BytesIO(excel_bayt({"Sözlük": tablo})))
    degerler = [c.value for r in wb["Sözlük"].iter_rows() for c in r]
    assert "=TOPLA(A1)" in degerler  # metin olarak yazıldı
    assert all(c.data_type != "f" for r in wb["Sözlük"].iter_rows() for c in r)


def test_sozluk_disa_aktar_yeniden_yuklenebilir():
    s = sozluk_olustur(["Sağlık bilimleri"])
    tablo = tm.sozluk_disa_aktar(s)
    yeni, _ = sozluk_tablosundan(tablo)
    assert [t.standart for t in yeni] == [t.standart for t in s]
    assert [t.varyantlar for t in yeni] == [t.varyantlar for t in s]


def test_alanlar_arasi_belirsiz():
    b = tm.alanlar_arasi_belirsiz(cekirdek_sozluk())
    assert "derin öğrenme modeli" in b["T0020"]  # eğitimdeki "derin öğrenme"
    assert any("lojistik regresyon modeli" in v for v in b.values())


# --- Ön işleme ile bütünleşme ------------------------------------------------------

METIN = ("Uzaktan eğitime katılan öğretmen adaylarının yapay zekâ okuryazarlığı ölçüldü. Ölçme ve "
         "değerlendirmede yapısal eşitlik modellemesinin kullanımı ve Tip 2 diyabetin yaygınlığı incelendi.")


def test_birim_temsili(E):
    b = on_isle(METIN, terimci=E)
    assert b.sozcukler[:4] == ["uzaktan_eğitim", "katılan", "öğretmen_adayları", "yapay_zekâ_okuryazarlığı"]
    assert "ölçme_ve_değerlendirme" in b.sozcukler and "tip_2_diyabet" in b.sozcukler
    assert [e.standart for e in b.terimler][:2] == ["uzaktan eğitim", "öğretmen adayları"]
    # Bölütler korunur: cümle sonu yine sınırdır
    assert len(b.bolutler) == 2


def test_terim_birimi_suzgeclerden_gecmez(E):
    lem = on_isle(METIN, OnIslemeAyarlari(lemmatizasyon=True), terimci=E)
    assert "uzaktan_eğitim" in lem.sozcukler and "uzak" not in lem.sozcukler
    uzun = on_isle(METIN, OnIslemeAyarlari(min_uzunluk=30), terimci=E)
    assert uzun.sozcukler == ["uzaktan_eğitim", "öğretmen_adayları", "yapay_zekâ_okuryazarlığı",
                              "ölçme_ve_değerlendirme", "yapısal_eşitlik_modellemesi", "tip_2_diyabet"]
    akad = on_isle(METIN, OnIslemeAyarlari(durak_sozcuk_akademik=True), terimci=E)
    assert "ölçme_ve_değerlendirme" in akad.sozcukler


def test_kullanici_durak_listesi_terim_birimini_cikarabilir(E):
    b = on_isle(METIN, OnIslemeAyarlari(ozel_durak_sozcukler=("uzaktan_eğitim",)), terimci=E)
    assert "uzaktan_eğitim" not in b.sozcukler and b.bolutler[0][0] is None


def test_kucuk_harf_ve_sapka_secenekleri_belirtece_uygulanir(E):
    assert "yapay_zeka_okuryazarlığı" in on_isle(METIN, OnIslemeAyarlari(sapkali_harf_normallestir=True),
                                                 terimci=E).sozcukler
    assert "Cronbach_alfa" in on_isle("Cronbach alfa katsayısı", OnIslemeAyarlari(kucuk_harf=False),
                                      terimci=E).sozcukler


def test_isaretle_temsili_ve_sozluksuz_davranis_ayni(E):
    isaret = TerimEslestirici(cekirdek_sozluk(), temsil="isaretle")
    for ayar in [OnIslemeAyarlari(), OnIslemeAyarlari(lemmatizasyon=True, durak_sozcuk_akademik=True)]:
        a = on_isle(METIN, ayar)
        b = on_isle(METIN, ayar, terimci=isaret)
        assert a.bolutler == b.bolutler and len(b.terimler) == 6 and a.terimler == []
        # Terim içermeyen metinde sözlük açıkken de çıktı birebir aynı
        m = "Bu metin; herhangi bir sözlük terimi içermiyor, 19. yüzyıl ve Prof. Dr. ifadeleriyle."
        assert on_isle(m, ayar).bolutler == on_isle(m, ayar, terimci=E).bolutler


def test_yer_tutucu_karakteri_metinde_olsa_bile(E):
    b = on_isle("ǂ0ǂ uzaktan eğitim", terimci=E)
    assert b.sozcukler == ["uzaktan_eğitim"]


def test_analizlerde_tek_birim(E):
    metinler = ["Uzaktan eğitimde öğrenci memnuniyeti.", "Uzaktan eğitimin sorunları ve uzaktan eğitime erişim.",
                "Yüz yüze eğitimde öğrenci memnuniyeti."]
    kapali = derlem_on_isle(metinler)
    acik = derlem_on_isle(metinler, terimci=E)
    f = sozcuk_sikligi(acik).set_index("Sözcük")["Sıklık"]
    assert f["uzaktan_eğitim"] == 3 and "uzaktan" not in f.index
    ng = ngram_sikligi(acik, 2).set_index("N-gram")["Sıklık"]
    assert ng["uzaktan_eğitim öğrenci"] == 1
    assert "uzaktan eğitimde" in set(ngram_sikligi(kapali, 2)["N-gram"])
    tf = tfidf_hesapla(acik)
    assert "uzaktan_eğitim" in tf.terimler
    ag = birliktelik_hesapla(acik, AgAyarlari(min_birlikte=1, aday_sozcuk=50, pencere=2))
    assert "uzaktan_eğitim" in ag.sozcukler


def test_tablolar(E):
    metinler = ["Uzaktan eğitime katılan öğrenciler.", "Yapay zekâsız üretken yapay zekâ olmaz."]
    b = derlem_on_isle(metinler, terimci=E)
    t = tm.eslesme_tablosu(b, ["A", "B"])
    assert list(t["Standart terim"]) == ["uzaktan eğitim", "üretken yapay zekâ"]
    for _, r in t.iterrows():
        m = metinler[["A", "B"].index(r["Belge"])]
        assert m[r["Başlangıç"]:r["Bitiş"]] == r["Metindeki biçim"]
    assert list(tm.cakisma_tablosu(b)["Çakışma türü"]) == ["iç içe"]
    assert list(tm.eslesmeyen_tablosu(b)["Neden"]) == ["türetim eki"]
    oz = tm.terim_ozeti(b)
    assert set(oz["Standart terim"]) == {"uzaktan eğitim", "üretken yapay zekâ"}
    tem = tm.temizlenmis_metinler(b, metinler, ["A", "B"])
    assert list(tem.columns) == ["Belge", "Satır", "Özgün metin", "Temizlenmiş metin", "Normalize terimler"]
    assert tem.loc[0, "Normalize terimler"] == "uzaktan eğitim"
    txt = tm.temizlenmis_txt(b, ["A", "B"]).decode("utf-8")
    assert txt.startswith("### A\nuzaktan_eğitim") and "### B" in txt


def test_ilk_n_karsilastir(E):
    metinler = ["Uzaktan eğitimde öğrenci memnuniyeti."] * 3 + ["Yüz yüze eğitimde öğrenci memnuniyeti."]
    kapali = sozcuk_sikligi(derlem_on_isle(metinler))
    acik = sozcuk_sikligi(derlem_on_isle(metinler, terimci=E))
    tablo, olcut = tm.ilk_n_karsilastir(kapali, acik, "Sözcük", "Sıklık", 5)
    assert olcut["acik_listede_terim_birimi"] == 1 and 0 <= olcut["jaccard"] <= 1
    assert "uzaktan_eğitim" in list(tablo["Sözlük açık"])


# --- Gizlilik: kullanıcı terimleri genel önbelleğe girmez --------------------------

def test_kullanici_terimi_modul_duzeyinde_saklanmaz():
    gizli = "zümrütkuşu yöntemi"
    E = TerimEslestirici([kullanici_terimi(gizli, terim_id="K1")])
    assert E.eslestir("zümrütkuşu yöntemine göre").eslesmeler
    on_isle("zümrütkuşu yöntemine göre", terimci=E)

    def ara(x, derinlik=0):
        if isinstance(x, str):
            return "zümrütkuşu" in x
        if derinlik > 3:
            return False
        if isinstance(x, dict):
            return any(ara(k, derinlik + 1) or ara(v, derinlik + 1) for k, v in x.items())
        if isinstance(x, (list, tuple, set, frozenset)):
            return any(ara(v, derinlik + 1) for v in x)
        return False

    import tmlab.preprocess as pp
    for modul in (tm, pp):
        assert not any(ara(v) for k, v in vars(modul).items() if not k.startswith("__"))
    # Ek önbelleğinin anahtarları yalnızca ses profilleridir
    assert tm._ek_kumesi.cache_info().currsize <= 64


# --- Bağımsız yapay test kümesi (yalnızca alt sınır denetimi) ----------------------

def _test_kumesi():
    yol = os.path.join(KOK, "degerlendirme", "veri", "terim_test_yapay.tsv")
    for satir in open(yol, encoding="utf-8"):
        if satir.startswith("#") or not satir.strip():
            continue
        cumle, beklenen = satir.rstrip("\n").split("\t")
        yield cumle, beklenen


def test_terim_test_kumesi_alt_sinir(E, EK):
    """Ayrıntılı ölçüm degerlendirme/v21_terim_dogrulama.py'dedir; burada gerilemeye karşı alt sınır denetlenir."""
    from degerlendirme.v21_terim_dogrulama import yapay_kumeyi_degerlendir

    s = yapay_kumeyi_degerlendir()
    assert s["genel"]["f1"] >= 0.95, s["genel"]
    assert s["genel"]["yanlis_pozitif"] <= 2, s["genel"]
