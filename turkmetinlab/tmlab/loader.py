"""CSV, Excel ve TXT dosyalarını okuma.

Yüklenen dosyalar yalnızca bellekte (bayt olarak) işlenir; diske yazılmaz.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import List, Optional, Tuple

import pandas as pd

# Denenecek kodlamalar (sıra önemlidir: UTF-8 katıdır, cp1254 neredeyse her
# baytı çözebildiği için sona yakın denenir).
KODLAMALAR = ("utf-8-sig", "utf-8", "cp1254", "iso-8859-9", "utf-16")

# Yanlış kodlamayla okunmuş UTF-8 metinde tipik bozuk dizgeler (mojibake)
_MOJIBAKE_RE = re.compile(r"Ã[\u0080-¿]|Ä[\u0080-¿]|Å[\u0080-¿]|Ä±|ÅŸ|Ã¼|Ã§|Ã¶|ÄŸ")

TXT_MODLARI = {
    "dosya": "Her dosya bir belge",
    "paragraf": "Her paragraf bir belge (boş satırla ayrılmış)",
    "satir": "Her satır bir belge",
}


class VeriOkumaHatasi(Exception):
    """Kullanıcıya gösterilecek, anlaşılır Türkçe hata."""


# --- İçerik türü denetimi ---------------------------------------------------

# Dosya imzaları (ilk baytlar). Uzantı değiştirilmiş dosyaları yakalamak için.
_IMZALAR = (
    (b"%PDF", "PDF belgesi"),
    (b"PK\x03\x04", "sıkıştırılmış dosya (ör. .xlsx, .docx, .zip)"),
    (b"\xd0\xcf\x11\xe0", "eski Microsoft Office dosyası (.xls/.doc)"),
    (b"\x89PNG", "PNG görseli"),
    (b"\xff\xd8\xff", "JPEG görseli"),
    (b"GIF8", "GIF görseli"),
    (b"\x1f\x8b", "gzip arşivi"),
    (b"Rar!", "RAR arşivi"),
    (b"7z\xbc\xaf", "7z arşivi"),
    (b"{\\rtf", "RTF belgesi"),
)

# Excel (.xlsx) açılmadan önce sıkıştırılmış içeriğin güvenli boyutta olduğu denetlenir
EXCEL_EN_FAZLA_ACIK_BOYUT = 200 * 2**20  # 200 MB (açılmış hali); openpyxl bunun birkaç katı bellek kullanır
EXCEL_EN_FAZLA_SIKISTIRMA_ORANI = 200


def _imza_turu(veri: bytes) -> Optional[str]:
    for imza, ad in _IMZALAR:
        if veri.startswith(imza):
            return ad
    return None


def metin_dosyasi_denetle(veri: bytes, ad: str = "Dosya") -> None:
    """CSV/TXT olarak yüklenen dosyanın gerçekten düz metin olduğunu denetler."""
    if not veri:
        raise VeriOkumaHatasi(f"{ad} boş.")
    tur = _imza_turu(veri)
    if tur:
        raise VeriOkumaHatasi(f"{ad} düz metin değil; bir {tur} gibi görünüyor. Dosya uzantısı yanlış olabilir. "
                              "Excel dosyalarını .xlsx uzantısıyla yükleyin.")
    if veri[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return  # UTF-16 metinde NUL baytlar olağandır
    ornek = veri[:65536]
    if b"\x00" in ornek:
        raise VeriOkumaHatasi(f"{ad} düz metin değil (ikili veri içeriyor).")
    kontrol = sum(1 for b in ornek if b < 32 and b not in (9, 10, 13, 12))
    if kontrol > max(10, len(ornek) * 0.02):
        raise VeriOkumaHatasi(f"{ad} düz metin değil gibi görünüyor (çok sayıda denetim karakteri).")


def excel_dosyasi_denetle(veri: bytes) -> None:
    """.xlsx dosyasının gerçek bir Excel çalışma kitabı ve güvenli boyutta olduğunu denetler."""
    if not veri:
        raise VeriOkumaHatasi("Dosya boş.")
    if not veri.startswith(b"PK\x03\x04"):
        tur = _imza_turu(veri)
        if veri.startswith(b"\xd0\xcf\x11\xe0"):
            raise VeriOkumaHatasi("Bu dosya eski .xls biçiminde. Excel'de 'Farklı Kaydet > Excel Çalışma Kitabı "
                                  "(.xlsx)' ile dönüştürüp yeniden yükleyin.")
        raise VeriOkumaHatasi("Bu dosya bir Excel çalışma kitabı değil" + (f"; bir {tur} gibi görünüyor." if tur else "."))
    try:
        with zipfile.ZipFile(io.BytesIO(veri)) as z:
            bilgiler = z.infolist()
            adlar = {b.filename for b in bilgiler}
            acik = sum(b.file_size for b in bilgiler)
            sikisik = sum(b.compress_size for b in bilgiler) or 1
    except zipfile.BadZipFile as e:
        raise VeriOkumaHatasi("Excel dosyası bozuk ya da eksik yüklenmiş.") from e
    if "xl/workbook.xml" not in adlar:
        tur = "Word belgesi" if "word/document.xml" in adlar else "sıkıştırılmış dosya"
        raise VeriOkumaHatasi(f"Bu dosya bir Excel çalışma kitabı değil; bir {tur} gibi görünüyor.")
    if acik > EXCEL_EN_FAZLA_ACIK_BOYUT or (acik > 50 * 2**20 and acik / sikisik > EXCEL_EN_FAZLA_SIKISTIRMA_ORANI):
        raise VeriOkumaHatasi("Excel dosyasının açılmış boyutu çok büyük "
                              f"({acik / 2**20:.0f} MB). Güvenlik ve bellek sınırları nedeniyle işlenmedi.")


@dataclass
class OkumaSonucu:
    tablo: pd.DataFrame
    kodlama: Optional[str] = None
    ayirici: Optional[str] = None
    uyarilar: List[str] = field(default_factory=list)


def baytlari_coz(veri: bytes) -> Tuple[str, str]:
    """Baytları uygun kodlamayla metne çevirir. (metin, kodlama) döndürür."""
    if not veri:
        raise VeriOkumaHatasi("Dosya boş.")
    if veri[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return veri.decode("utf-16"), "utf-16"
    for kod in KODLAMALAR:
        if kod == "utf-16":
            continue
        try:
            return veri.decode(kod), kod
        except UnicodeDecodeError:
            continue
    raise VeriOkumaHatasi("Dosyanın karakter kodlaması tanınamadı. Lütfen dosyayı UTF-8 olarak kaydedin.")


def mojibake_var_mi(metin: str) -> bool:
    """Metinde 'Ã¼', 'ÅŸ' gibi bozuk Türkçe karakter izleri var mı?"""
    return bool(_MOJIBAKE_RE.search(metin[:200000]))


def _ayirici_bul(ornek: str) -> str:
    try:
        return csv.Sniffer().sniff(ornek, delimiters=",;\t|").delimiter
    except csv.Error:
        # Basit sayım yedeği
        ilk_satir = ornek.splitlines()[0] if ornek else ""
        adaylar = {a: ilk_satir.count(a) for a in [";", ",", "\t", "|"]}
        en_iyi = max(adaylar, key=adaylar.get)
        return en_iyi if adaylar[en_iyi] > 0 else ","


def csv_oku(veri: bytes, baslik_satiri: bool = True) -> OkumaSonucu:
    metin_dosyasi_denetle(veri, "CSV dosyası")
    metin, kodlama = baytlari_coz(veri)
    ayirici = _ayirici_bul(metin[:20000])
    # Ayrıştırma csv modülüyle açıkça yapılır: pandas, başlıktan fazla alan
    # içeren satırlarda sütunları sessizce kaydırabildiği için kullanılmaz.
    csv.field_size_limit(2**31 - 1)
    try:
        satirlar = list(csv.reader(io.StringIO(metin, newline=""), delimiter=ayirici))
    except csv.Error as e:
        raise VeriOkumaHatasi(f"CSV dosyası okunamadı: {e}") from e
    satirlar = [s for s in satirlar if any(h.strip() for h in s)]
    if not satirlar:
        raise VeriOkumaHatasi("Dosyada veri bulunamadı.")
    if baslik_satiri:
        baslik, veri_satirlari, ilk_no = satirlar[0], satirlar[1:], 2
    else:
        genislik = max(len(s) for s in satirlar)
        baslik, veri_satirlari, ilk_no = [f"Sütun {i + 1}" for i in range(genislik)], satirlar, 1
    genislik = len(baslik)
    gecerli, hatali, eksik = [], [], 0
    for no, s in enumerate(veri_satirlari, start=ilk_no):
        if len(s) > genislik and any(h.strip() for h in s[genislik:]):
            hatali.append(no)
            continue
        if len(s) < genislik and any(str(b).strip() for b in baslik[len(s):]):
            eksik += 1
        s = (s + [""] * genislik)[:genislik]
        gecerli.append(s)
    tablo = pd.DataFrame(gecerli, columns=baslik, dtype=object)
    tablo = tablo.replace({"": None})
    uyarilar = []
    if eksik:
        uyarilar.append(f"{eksik} satırda başlıktan az sütun vardı; eksik hücreler boş kabul edildi.")
    supheli = _kapanmamis_tirnak(gecerli, ayirici, ilk_no)
    if supheli:
        uyarilar.append(f"Satır {supheli} çevresinde kapanmamış bir tırnak işareti olabilir: bu hücre, dosyanın "
                        "birçok satırını tek hücrede birleştirmiş görünüyor. Dosyayı kontrol edin.")
    if hatali:
        ornek = ", ".join(map(str, hatali[:10])) + (" …" if len(hatali) > 10 else "")
        uyarilar.append(f"{len(hatali)} satır, başlıktan fazla sütun içerdiği için analize alınmadı "
                        f"(satır no: {ornek}). Metin içinde ayırıcı karakter varsa hücrenin tırnak içinde "
                        "olduğundan emin olun.")
    if mojibake_var_mi(metin):
        uyarilar.append("Dosyada bozuk Türkçe karakterler (ör. 'Ã¼', 'ÅŸ') görünüyor. Dosya daha önce "
                        "yanlış kodlamayla kaydedilmiş olabilir; sonuçları kontrol edin.")
    return OkumaSonucu(_temizle(tablo), kodlama, ayirici, uyarilar)


def _kapanmamis_tirnak(satirlar, ayirici: str, ilk_no: int) -> Optional[int]:
    """Kapanmamış bir tırnak, sonraki satırları tek hücreye yutar. Böyle bir hücre arar."""
    for i, s in enumerate(satirlar):
        for h in s:
            if h and h.count("\n") >= 10:
                alt = h.split("\n")
                ayiricili = sum(1 for a in alt if a.count(ayirici) >= 1)
                if ayiricili >= 0.3 * len(alt):
                    return i + ilk_no
    return None


def excel_sayfalari(veri: bytes) -> List[str]:
    excel_dosyasi_denetle(veri)
    try:
        return list(pd.ExcelFile(io.BytesIO(veri), engine="openpyxl").sheet_names)
    except Exception as e:
        raise VeriOkumaHatasi("Excel dosyası açılamadı. Yalnızca .xlsx biçimi desteklenir "
                              "(eski .xls dosyalarını Excel'de 'Farklı Kaydet > .xlsx' ile dönüştürün).") from e


def excel_oku(veri: bytes, sayfa=0, baslik_satiri: bool = True) -> OkumaSonucu:
    excel_dosyasi_denetle(veri)
    try:
        # keep_default_na=False: "NA", "null" gibi değerler sessizce boşa çevrilmez;
        # yer tutucular daha sonra açıkça raporlanarak ayıklanır.
        tablo = pd.read_excel(io.BytesIO(veri), sheet_name=sayfa, dtype=str, engine="openpyxl",
                              header=0 if baslik_satiri else None, keep_default_na=False, na_values=[""])
    except Exception as e:
        raise VeriOkumaHatasi("Excel dosyası okunamadı. Yalnızca .xlsx biçimi desteklenir.") from e
    if not baslik_satiri:
        tablo.columns = [f"Sütun {i + 1}" for i in range(len(tablo.columns))]
    uyarilar = []
    formul = excel_formul_sayisi(veri, sayfa)
    if formul:
        uyarilar.append(f"Bu sayfada {formul} hücre formül içeriyor. Formüller çalıştırılmadı; Excel'in dosyaya "
                        "kaydettiği son hesaplanmış değerler kullanıldı (dosya Excel dışında üretildiyse bu "
                        "hücreler boş görünebilir).")
    return OkumaSonucu(_temizle(tablo), uyarilar=uyarilar)


def excel_formul_sayisi(veri: bytes, sayfa=0) -> int:
    """Seçili sayfadaki formül hücrelerini sayar (formüller hiçbir zaman çalıştırılmaz)."""
    from openpyxl import load_workbook

    try:
        kitap = load_workbook(io.BytesIO(veri), read_only=True, data_only=False)
        ws = kitap.worksheets[sayfa] if isinstance(sayfa, int) else kitap[sayfa]
        sayi = 0
        for satir in ws.iter_rows():
            for h in satir:
                if getattr(h, "data_type", None) == "f":
                    sayi += 1
        kitap.close()
        return sayi
    except Exception:
        return 0


def txt_belgelere_ayir(metin: str, mod: str = "dosya") -> List[str]:
    if mod not in TXT_MODLARI:
        raise ValueError(f"Geçersiz TXT modu: {mod}")
    metin = metin.replace("\r\n", "\n").replace("\r", "\n")
    if mod == "dosya":
        parcalar = [metin]
    elif mod == "paragraf":
        parcalar = re.split(r"\n\s*\n", metin)
    else:
        parcalar = metin.split("\n")
    return [p.strip() for p in parcalar if p and p.strip()]


def txt_oku(dosyalar: List[Tuple[str, bytes]], mod: str = "dosya") -> OkumaSonucu:
    """Bir veya daha fazla TXT dosyasını belge tablosuna çevirir.

    dosyalar: (dosya_adı, bayt_içerik) listesi
    """
    satirlar, uyarilar, kodlamalar = [], [], set()
    for ad, veri in dosyalar:
        try:
            metin_dosyasi_denetle(veri, f"'{ad}'")
            metin, kod = baytlari_coz(veri)
        except VeriOkumaHatasi as e:
            uyarilar.append(f"{e} Bu dosya atlandı." if veri else f"'{ad}' boş; atlandı.")
            continue
        kodlamalar.add(kod)
        if mojibake_var_mi(metin):
            uyarilar.append(f"'{ad}' dosyasında bozuk Türkçe karakterler görünüyor.")
        parcalar = txt_belgelere_ayir(metin, mod)
        for i, p in enumerate(parcalar, start=1):
            etiket = ad if mod == "dosya" else f"{ad} — {'paragraf' if mod == 'paragraf' else 'satır'} {i}"
            satirlar.append({"belge": etiket, "metin": p})
    if not satirlar:
        raise VeriOkumaHatasi("Yüklenen TXT dosyalarında okunabilir metin bulunamadı.")
    return OkumaSonucu(pd.DataFrame(satirlar), ", ".join(sorted(kodlamalar)), None, uyarilar)



DOCX_MODLARI = {
    "dosya": "Her Word dosyası bir belge",
    "paragraf": "Her paragraf / tablo satırı bir belge",
}
DOCX_EN_FAZLA_DOSYA = 16 * 2**20
DOCX_EN_FAZLA_ACIK_BOYUT = 80 * 2**20


def docx_dosyasi_denetle(veri: bytes) -> None:
    """Office Open XML yapısını ve ZIP bombası riskini denetle."""
    if not veri:
        raise VeriOkumaHatasi("Word dosyası boş.")
    if len(veri) > DOCX_EN_FAZLA_DOSYA:
        raise VeriOkumaHatasi("Word dosyası 16 MB sınırını aşıyor.")
    if not veri.startswith(b"PK\x03\x04"):
        raise VeriOkumaHatasi("Bu dosya geçerli bir .docx Word belgesi değil.")
    try:
        with zipfile.ZipFile(io.BytesIO(veri)) as z:
            bilgiler = z.infolist()
            adlar = {b.filename for b in bilgiler}
            toplam = sum(b.file_size for b in bilgiler)
            sikisik = sum(b.compress_size for b in bilgiler) or 1
    except (zipfile.BadZipFile, OSError) as e:
        raise VeriOkumaHatasi("Word dosyası bozuk veya eksik.") from e
    if "word/document.xml" not in adlar or "[Content_Types].xml" not in adlar:
        raise VeriOkumaHatasi("Bu dosya bir .docx Word belgesi değil.")
    if "word/vbaProject.bin" in adlar:
        raise VeriOkumaHatasi("Makro içeren Word dosyaları desteklenmiyor.")
    if toplam > DOCX_EN_FAZLA_ACIK_BOYUT or (toplam > 10 * 2**20 and toplam / sikisik > 200):
        raise VeriOkumaHatasi("Word dosyasının açılmış boyutu güvenli sınırı aşıyor.")


def docx_oku(dosyalar: List[Tuple[str, bytes]], mod: str = "dosya") -> OkumaSonucu:
    """Word paragraf ve tablo satırlarını belge sırasıyla, bellekte okur."""
    if mod not in DOCX_MODLARI:
        raise VeriOkumaHatasi("Geçersiz Word bölme seçeneği.")
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    satirlar, uyarilar = [], []
    for ad, veri in dosyalar:
        docx_dosyasi_denetle(veri)
        try:
            belge = Document(io.BytesIO(veri))
            parcalar = []
            for blok in belge.iter_inner_content():
                if isinstance(blok, Paragraph):
                    metin = blok.text.strip()
                    if metin:
                        parcalar.append(metin)
                elif isinstance(blok, Table):
                    for satir in blok.rows:
                        alanlar = [h.text.strip() for h in satir.cells]
                        if any(alanlar):
                            parcalar.append(" | ".join(alanlar))
        except Exception as e:
            raise VeriOkumaHatasi("Word dosyası okunamadı. Parolalı veya bozuk olabilir.") from e
        if not parcalar:
            uyarilar.append(f"'{ad}' içinde okunabilir paragraf veya tablo metni bulunamadı.")
            continue
        if mod == "dosya":
            satirlar.append({"belge": ad, "metin": "\n\n".join(parcalar)})
        else:
            for i, metin in enumerate(parcalar, start=1):
                satirlar.append({"belge": f"{ad} — bölüm {i}", "metin": metin})
    if not satirlar:
        raise VeriOkumaHatasi("Word dosyalarında okunabilir metin bulunamadı.")
    uyarilar.append("Resimlerdeki metinler, üstbilgiler, altbilgiler ve dipnotlar çıkarılmadı.")
    return OkumaSonucu(pd.DataFrame(satirlar), uyarilar=uyarilar)


def _temizle(tablo: pd.DataFrame) -> pd.DataFrame:
    tablo = tablo.dropna(how="all").dropna(axis=1, how="all")
    tablo.columns = [str(c).strip() if str(c).strip() else f"Sütun {i + 1}" for i, c in enumerate(tablo.columns)]
    # Yinelenen sütun adlarını tekilleştir
    gorulen = {}
    yeni = []
    for c in tablo.columns:
        if c in gorulen:
            gorulen[c] += 1
            yeni.append(f"{c} ({gorulen[c]})")
        else:
            gorulen[c] = 0
            yeni.append(c)
    tablo.columns = yeni
    if tablo.empty:
        raise VeriOkumaHatasi("Dosyada veri bulunamadı.")
    return tablo.reset_index(drop=True)


def metin_sutunu_oner(tablo: pd.DataFrame) -> Optional[str]:
    """Ortalama uzunluğu en fazla olan metin sütununu önerir."""
    en_iyi, en_uzun = None, -1.0
    for c in tablo.columns:
        degerler = tablo[c].dropna().astype(str)
        if degerler.empty:
            continue
        # Çoğunluğu sayı olan sütunları atla
        sayisal_oran = degerler.str.fullmatch(r"\s*[-+]?\d+([.,]\d+)?\s*").mean()
        if sayisal_oran > 0.8:
            continue
        ort = degerler.str.len().mean()
        if ort > en_uzun:
            en_iyi, en_uzun = c, ort
    return en_iyi


def dosya_oku(ad: str, veri: bytes, sayfa=0, txt_modu: str = "dosya", baslik_satiri: bool = True) -> OkumaSonucu:
    """Uzantıya göre uygun okuyucuyu çağırır."""
    uzanti = ad.lower().rsplit(".", 1)[-1] if "." in ad else ""
    if uzanti in ("csv", "tsv"):
        return csv_oku(veri, baslik_satiri)
    if uzanti in ("xlsx", "xlsm"):
        return excel_oku(veri, sayfa, baslik_satiri)
    if uzanti == "xls":
        raise VeriOkumaHatasi("Eski .xls biçimi desteklenmiyor. Excel'de 'Farklı Kaydet > Excel Çalışma "
                              "Kitabı (.xlsx)' seçeneğiyle dönüştürüp yeniden yükleyin.")
    if uzanti == "txt":
        return txt_oku([(ad, veri)], txt_modu)
    if uzanti == "docx":
        return docx_oku([(ad, veri)], txt_modu)
    raise VeriOkumaHatasi(f"Desteklenmeyen dosya türü: .{uzanti}. CSV, XLSX, TXT veya DOCX yükleyin.")
